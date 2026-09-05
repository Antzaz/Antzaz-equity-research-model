from __future__ import annotations

"""Run the AI-growth forecasting layer for one equity-research workbook.

Examples:
    python ai_growth_research.py GOOGL
    python ai_growth_research.py GOOGL --llm
    python ai_growth_research.py GOOGL --llm --model gpt-5.6-terra

The engine never changes the authoritative DCF assumptions. It writes a separate
"AI Growth Forecast" research sheet and a JSON audit artifact. Business-model policy gates
industrial FCF/reverse-DCF outputs for banks, insurers, broker/dealers and REITs.
"""

import argparse
from datetime import datetime
import json
from pathlib import Path
import re

from openpyxl import load_workbook

from business_model_registry import workbook_policy, reverse_dcf_applicability_message
from chart_excel_compat import apply_chart_compatibility_fix
from machine_learning.ai_growth import (
    LightGBMGrowthForecaster,
    ai_adjustments,
    apply_ai_overlay,
    deterministic_ai_signals,
    expectations_gap,
    growth_training_frame_from_history,
    llm_ai_signals,
    load_kpi_evidence,
    reverse_dcf_from_workbook,
    workbook_current_growth_features,
)
from machine_learning.ai_growth_output import write_ai_growth_sheet
from machine_learning.data import latest_workbook

BASE = Path(__file__).resolve().parent
DEFAULT_DB = BASE / "ml_data" / "ml_history.sqlite"
RUNS = BASE / "ml_runs"

PLACEHOLDER_TERMS=(
    "analyst input","to be updated","placeholder","not available","not disclosed",
    "tbd","n/a","n/m","unknown",
)
AI_CONTEXT_TERMS=(
    " ai ","artificial intelligence","generative ai","genai","machine learning","llm",
    "large language model","inference","training compute","gpu","accelerator","agentic",
    "copilot","foundation model","ai-","ai ",
)


def ticker_type(raw: str) -> str:
    ticker = str(raw or "").upper().strip()
    if not re.fullmatch(r"[A-Z0-9.\-]{1,10}", ticker):
        raise argparse.ArgumentTypeError("Enter a ticker such as GOOGL, MSFT, NVDA, TSM, or SIE.DE")
    return ticker


def parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="LLM evidence extraction + LightGBM fundamental growth + sector-safe valuation expectations"
    )
    p.add_argument("ticker", type=ticker_type)
    p.add_argument("--workbook", help="Workbook path; defaults to the newest generated model")
    p.add_argument("--history-db", default=str(DEFAULT_DB), help="Persistent point-in-time ML SQLite database")
    p.add_argument("--llm", action="store_true", help="Use OpenAI structured extraction when OPENAI_API_KEY is set")
    p.add_argument("--model", default="gpt-5.6-luna", help="OpenAI extraction model override")
    p.add_argument("--no-workbook-write", action="store_true")
    return p


def _policy_for_workbook(path: Path, ticker: str):
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
        return workbook_policy(wb, ticker)
    except Exception:
        from business_model_registry import get_business_model_policy
        return get_business_model_policy(ticker)


def _row_text(row: dict) -> str:
    return " | ".join(
        str(row.get(k) or "")
        for k in ("kpi","name","current","value","unit_comparison","signal","investment_read_through","data_type")
    ).strip()


def _is_substantive_ai_row(row: dict) -> bool:
    """Keep real AI evidence; reject analyst-input/template rows that create false precision."""
    text=_row_text(row)
    low=f" {text.lower()} "
    if not any(term in low for term in AI_CONTEXT_TERMS):
        return False
    value_text=" | ".join(
        str(row.get(k) or "")
        for k in ("current","value","unit_comparison","signal","investment_read_through","data_type")
    ).strip()
    value_low=value_text.lower().strip()
    if not value_low:
        return False
    numeric=bool(re.search(r"[-+]?\d+(?:\.\d+)?",value_text))
    substantive_words=any(
        word in value_low for word in (
            "company reported","management","guidance","grew","growth","revenue","customer",
            "user","backlog","bookings","margin","capex","contract","deployment","usage","paid",
        )
    )
    placeholder_only=any(term in value_low for term in PLACEHOLDER_TERMS) and not numeric and not substantive_words
    return not placeholder_only


def _clean_ai_evidence(rows: list[dict], corpus: str) -> tuple[list[dict],str,dict]:
    """Remove template placeholders before deterministic or LLM scoring.

    Source-file evidence remains available, but placeholder lines from KPI templates are removed.
    This prevents labels such as 'AI revenue — Analyst input' from nudging neutral scores upward.
    """
    clean_rows=[dict(r) for r in rows if isinstance(r,dict) and _is_substantive_ai_row(r)]
    clean_lines=[]
    removed_lines=0
    for line in str(corpus or "").splitlines():
        low=line.lower().strip()
        if low and any(term in low for term in PLACEHOLDER_TERMS) and not re.search(r"[-+]?\d+(?:\.\d+)?",line):
            removed_lines+=1
            continue
        clean_lines.append(line)
    clean_corpus="\n".join(clean_lines)
    stats={
        "raw_kpi_rows":len(rows),
        "substantive_ai_rows":len(clean_rows),
        "filtered_placeholder_rows":max(0,len(rows)-len(clean_rows)),
        "filtered_placeholder_lines":removed_lines,
    }
    return clean_rows,clean_corpus,stats


def main() -> int:
    args = parser().parse_args()
    ticker = args.ticker
    workbook = Path(args.workbook).resolve() if args.workbook else latest_workbook(BASE, ticker)
    history_db = Path(args.history_db).resolve()
    policy = _policy_for_workbook(workbook, ticker)

    raw_rows, raw_corpus = load_kpi_evidence(BASE, ticker)
    rows, corpus, evidence_quality = _clean_ai_evidence(raw_rows, raw_corpus)
    signals = (
        llm_ai_signals(rows, corpus, model=args.model)
        if args.llm
        else deterministic_ai_signals(rows, corpus)
    )
    if evidence_quality["substantive_ai_rows"]==0 and not corpus.strip():
        signals.summary="No substantive company-specific AI evidence was available after placeholder filtering; scores remain neutral."
    print(
        f"[ai-growth] evidence: {signals.evidence_count} substantive KPI row(s) "
        f"({evidence_quality['filtered_placeholder_rows']} filtered); "
        f"extraction={signals.extraction_mode}; confidence={signals.confidence:.0%}"
    )
    print(f"[ai-growth] business model: {policy.label}; primary valuation={policy.primary_valuation}")

    frame = growth_training_frame_from_history(history_db)
    current = workbook_current_growth_features(workbook, ticker)
    model = LightGBMGrowthForecaster()
    revenue = model.fit_predict(
        frame,
        current,
        target_col="target_next_revenue_growth",
        target_label="Next FY revenue growth",
    )
    fcf = model.fit_predict(
        frame,
        current,
        target_col="target_next_fcf_growth",
        target_label="Next FY FCF growth",
    )

    adjustments = ai_adjustments(signals)
    ai_revenue = apply_ai_overlay(revenue, adjustments["revenue_growth_adjustment"])

    if policy.industrial_fcf_primary:
        ai_fcf = apply_ai_overlay(fcf, adjustments["fcf_growth_adjustment"])
    else:
        fcf.status = "NOT_APPLICABLE"
        fcf.prediction = None
        fcf.confidence = "N/M"
        fcf.metrics = {
            "business_model_gate": True,
            "reason": f"Industrial FCF growth is not a primary research target for {policy.label}.",
            "primary_valuation": policy.primary_valuation,
        }
        fcf.drivers = []
        ai_fcf = None
        adjustments["fcf_growth_adjustment"] = 0.0

    if policy.reverse_dcf_allowed:
        reverse_dcf = reverse_dcf_from_workbook(workbook)
        gap = expectations_gap(ai_fcf, reverse_dcf)
    else:
        reverse_dcf = {
            "status": "NOT_APPLICABLE",
            "implied_annual_fcf_growth": None,
            "business_model": policy.key,
            "primary_valuation": policy.primary_valuation,
            "reason": reverse_dcf_applicability_message(policy),
        }
        gap = {
            "status": "NOT_APPLICABLE",
            "fcf_growth_gap": None,
            "interpretation": reverse_dcf_applicability_message(policy),
        }

    payload = {
        "ticker": ticker,
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "workbook": str(workbook),
        "history_db": str(history_db),
        "business_model_policy": policy.to_dict(),
        "evidence_quality": evidence_quality,
        "architecture": {
            "evidence_layer": signals.extraction_mode,
            "forecast_model": "LightGBM",
            "baseline_model": "ElasticNet",
            "explainability": "SHAP with LightGBM feature-importance fallback",
            "valuation_bridge": "reverse DCF implied FCF growth" if policy.reverse_dcf_allowed else policy.primary_valuation,
            "governance": (
                "AI evidence is a bounded overlay until sufficient dated AI KPI history exists "
                "to train AI features directly. Template/analyst-input rows are filtered before scoring. "
                "This layer does not overwrite DCF assumptions. Business-model policy disables industrial "
                "FCF/reverse-DCF where that framework is not economically appropriate."
            ),
        },
        "ai_signals": signals.to_dict(),
        "revenue_forecast": revenue.to_dict(),
        "fcf_forecast": fcf.to_dict(),
        "ai_adjustments": adjustments,
        "ai_adjusted_revenue_growth": ai_revenue,
        "ai_adjusted_fcf_growth": ai_fcf,
        "reverse_dcf": reverse_dcf,
        "expectations_gap": gap,
    }

    stamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M%S")
    run_dir = RUNS / ticker / stamp
    run_dir.mkdir(parents=True, exist_ok=True)
    output = run_dir / "ai_growth_results.json"
    output.write_text(json.dumps(payload, indent=2, ensure_ascii=False, default=str) + "\n", encoding="utf-8")

    if not args.no_workbook_write:
        write_ai_growth_sheet(workbook, ticker, payload)
        compat = apply_chart_compatibility_fix(workbook)
        print(f"[ai-growth] Excel chart compatibility: {compat.get('charts_repaired',0)} chart(s) repaired")

    def show(label: str, value):
        print(f"[ai-growth] {label}: {'N/M' if value is None else f'{value:.1%}'}")

    show("LightGBM revenue growth", revenue.prediction)
    show("AI-adjusted revenue growth", ai_revenue)
    show("LightGBM FCF growth", fcf.prediction)
    show("AI-adjusted FCF growth", ai_fcf)
    show("reverse-DCF implied FCF growth", reverse_dcf.get("implied_annual_fcf_growth"))
    show("AI expectations gap", gap.get("fcf_growth_gap"))
    print(f"[ai-growth] output: {output}")
    if not args.no_workbook_write:
        print(f"[ai-growth] workbook sheet updated: {workbook} -> AI Growth Forecast")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
