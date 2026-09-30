from openpyxl import Workbook

from supply_chain_intelligence import _candidate_sources
from segment_analysis_v2 import _valid_segment_name
from numeric_coverage_audit import ensure_numeric_coverage_audit


def test_segment_name_filter_rejects_broken_fragments():
    assert _valid_segment_name("Compute &amp; Networking")=="Compute & Networking"
    assert _valid_segment_name('Networking" and "Graphics') is None
    assert _valid_segment_name("Google Cloud")=="Google Cloud"


def test_supplier_sources_accept_direct_regulatory_fallback(monkeypatch):
    wb=Workbook()
    monkeypatch.setattr("supply_chain_intelligence._latest_regulatory_filing",lambda ticker:"https://sec.example/latest.htm")
    monkeypatch.setattr("supply_chain_intelligence.issuer_sources",lambda ticker,website=None:{})
    sources=_candidate_sources(wb,"ANY",{})
    assert sources[0]==("https://sec.example/latest.htm","Regulatory / annual filing")


def test_numeric_coverage_audit_flags_styled_blank_numeric_field():
    wb=Workbook(); ws=wb.active; ws.title="Capital Allocation"
    ws["A1"]="Metric"; ws["B1"]="Value"
    ws["A2"]="Invested Capital"; ws["B2"]=None; ws["B2"].number_format='#,##0.0'
    out=ensure_numeric_coverage_audit(wb,"TEST")
    assert out["blank_numeric_fields"]>=1
    assert out["review_fields"]>=1
    assert "Numeric Coverage Audit" in wb.sheetnames
