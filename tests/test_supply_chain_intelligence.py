from openpyxl import Workbook

from supply_chain_intelligence import SupplierEvidence, ensure_supply_chain_intelligence


def test_supply_chain_sheet_uses_verified_records_and_risk_summary():
    wb=Workbook()
    recs=[
        SupplierEvidence(
            supplier="Example Foundry",
            relationship="Single / sole-source supplier",
            evidence_status="Disclosed",
            importance_score=95,
            dependency_score=90,
            risk_flags="Single / sole source; Geopolitical / geographic",
            evidence="The company identifies Example Foundry as a sole source manufacturing supplier.",
            source="https://example.com/filing",
            source_type="Regulatory / annual filing",
        ),
        SupplierEvidence(
            supplier="Example Components",
            relationship="Supply agreement",
            evidence_status="Confirmed public evidence",
            importance_score=65,
            dependency_score=45,
            risk_flags="No specific risk flag in retrieved evidence",
            evidence="The company has a supply agreement with Example Components.",
            source="https://example.com/ir",
            source_type="Issuer-owned",
        ),
    ]
    result=ensure_supply_chain_intelligence(wb,"TEST",records=recs)

    assert result["status"]=="PUBLIC-EVIDENCE"
    assert result["supplier_count"]==2
    assert result["high_dependency_count"]==1
    assert result["single_source_count"]==1
    assert result["geopolitical_flag_count"]==1
    ws=wb["Supply Chain"]
    assert ws["B6"].value=="Example Foundry"
    assert ws["D6"].value=="Disclosed"
    assert len(ws._charts)==1


def test_supply_chain_sheet_keeps_missing_evidence_as_review():
    wb=Workbook()
    result=ensure_supply_chain_intelligence(wb,"TEST",records=[])

    assert result["status"]=="REVIEW"
    assert result["supplier_count"]==0
    assert "No named supplier relationships were verified" in wb["Supply Chain"]["A6"].value
    assert len(wb["Supply Chain"]._charts)==0
