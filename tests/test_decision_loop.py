from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook

from institutional_research.src import decision_loop as dl


def test_realized_portfolio_twr_and_point_in_time_weights():
    ledger=pd.DataFrame([
        {"Date":"2026-01-02","Ticker":"","Action":"DEPOSIT","CashFlow":1000},
        {"Date":"2026-01-02","Ticker":"AAA","Action":"BUY","Shares":10,"Price":100,"Fees":0},
    ])
    ledger["Date"]=pd.to_datetime(ledger["Date"])
    for c in dl.LEDGER_COLUMNS:
        if c not in ledger: ledger[c]=np.nan
    ledger=ledger[dl.LEDGER_COLUMNS]
    idx=pd.to_datetime(["2026-01-02","2026-01-05","2026-01-06"])
    raw=pd.DataFrame({"AAA":[100.0,110.0,105.0]},index=idx)
    bench=pd.Series([100.0,102.0,103.0],index=idx)
    result=dl.build_realized_portfolio(ledger,raw,bench,{})
    assert result["summary"]["status"]=="PASS"
    assert abs(result["summary"]["twr_total"]-0.05)<1e-9
    assert result["summary"]["mwr_xirr"] is not None
    weights=result["weights"]
    last=weights[(weights["Date"]==idx[-1])&(weights["Ticker"]=="AAA")].iloc[0]
    assert abs(last["Weight"]-1.0)<1e-9


def test_realized_portfolio_dividend_is_attributed():
    ledger=pd.DataFrame([
        {"Date":"2026-01-02","Ticker":"","Action":"DEPOSIT","CashFlow":1000},
        {"Date":"2026-01-02","Ticker":"AAA","Action":"BUY","Shares":10,"Price":100,"Fees":0},
    ])
    ledger["Date"]=pd.to_datetime(ledger["Date"])
    for c in dl.LEDGER_COLUMNS:
        if c not in ledger: ledger[c]=np.nan
    ledger=ledger[dl.LEDGER_COLUMNS]
    idx=pd.to_datetime(["2026-01-02","2026-01-05"])
    raw=pd.DataFrame({"AAA":[100.0,100.0]},index=idx)
    actions={"AAA":pd.DataFrame({"Dividends":[0.0,1.0],"Stock Splits":[0.0,0.0]},index=idx)}
    result=dl.build_realized_portfolio(ledger,raw,None,actions)
    assert abs(result["summary"]["twr_total"]-0.01)<1e-9
    attr=result["attribution"]
    assert abs(attr[attr["Ticker"]=="AAA"]["DividendContribution"].sum()-0.01)<1e-9


def test_decision_journal_scores_buy_and_trim(tmp_path):
    path=tmp_path/"journal.csv"
    pd.DataFrame([
        {"Date":"2025-01-02","Ticker":"AAA","Decision":"BUY","Conviction":5},
        {"Date":"2025-01-02","Ticker":"BBB","Decision":"TRIM","Conviction":4},
    ]).to_csv(path,index=False)
    idx=pd.bdate_range("2025-01-02","2026-02-01")
    n=len(idx)
    prices=pd.DataFrame({
        "AAA":100*np.linspace(1,1.30,n),
        "BBB":100*np.linspace(1,0.90,n),
        "SPY":100*np.linspace(1,1.10,n),
    },index=idx)
    detail,summary=dl.decision_journal_analytics(path,prices,"SPY")
    matured=detail[detail["Matured"]==True]
    assert not matured.empty
    assert matured["Correct"].all()
    assert not summary.empty


def _research_book(path):
    wb=Workbook()
    c=wb.active;c.title="Company Data";c["B8"]=100;c["B9"]=1
    d=wb.create_sheet("Decision View")
    rows=[
        ("Current Market Price",100),
        ("Base DCF Fair Value",133.1),
        ("MODEL VIEW","ATTRACTIVE"),
        ("Overall Investment Score",75),
        ("Business Quality",80),
        ("Valuation / Price",65),
    ]
    for i,(k,v) in enumerate(rows,1): d.cell(i,1,k);d.cell(i,2,v)
    a=wb.create_sheet("Advanced Analytics")
    a["A1"]="P10 Value / Share";a["B1"]=80
    a["A2"]="P90 Value / Share";a["B2"]=170
    q=wb.create_sheet("Data Quality")
    q.append(["Control","PASS"]);q.append(["Control2","PASS"]);q.append(["Control3","REVIEW"])
    f=wb.create_sheet("Forecast Accountability")
    f.append(["","",""])
    for r,(year,val) in enumerate([(2027,120),(2030,160)],7):
        f.cell(r,1,year);f.cell(r,2,"Revenue");f.cell(r,3,val)
    ca=wb.create_sheet("Capital Allocation")
    ca["A1"]="Latest net share reduction";ca["B1"]=0.02
    wb.save(path)


def test_research_bridge_and_sizing(tmp_path):
    path=tmp_path/"AAA_Equity_Research_20260923_120000.xlsx"
    _research_book(path)
    bridge=dl.research_expected_return_bridge(
        ["AAA"],tmp_path,{"AAA":{"dividendYield":0.01}},
        benchmark_expected_return=0.08,convergence_years=3,
    )
    assert len(bridge)==1
    row=bridge.iloc[0]
    assert row["SourceStatus"]=="PASS"
    assert row["RawExpectedReturn"]>0.10
    assert row["ConfidenceAdjustedExpectedReturn"]<row["RawExpectedReturn"]

    holdings=pd.DataFrame([{"Ticker":"AAA","Weight":0.15,"RiskContributionPct":0.18,"MarketValue":15000}])
    sizing=dl.position_sizing_ranges(holdings,bridge,max_position=0.25)
    assert len(sizing)==1
    assert 0<=sizing.iloc[0]["SuggestedMin"]<=sizing.iloc[0]["SuggestedMax"]<=0.25
    budget=dl.thesis_budget(holdings,bridge)
    assert "ExpectedAlphaBudget" in budget


def test_rebalance_cost_gate():
    weights=pd.DataFrame([
        {"Portfolio":"Regime-Aware ML Maximum Sharpe","Ticker":"AAA","CurrentWeight":0.10,"TargetWeight":0.20,"WeightChange":0.10}
    ])
    holdings=pd.DataFrame([{"Ticker":"AAA","Weight":0.10,"MarketValue":10000}])
    liquidity=pd.DataFrame([{"Ticker":"AAA","AverageDailyDollarVolume":1_000_000}])
    bridge=pd.DataFrame([{"Ticker":"AAA","ExpectedAlpha":0.10}])
    out=dl.transaction_cost_rebalance(weights,holdings,liquidity,bridge,rebalance_threshold=.03)
    assert len(out)==1
    assert out.iloc[0]["EstimatedCostBps"]>0
    assert out.iloc[0]["Decision"] in {"REBALANCE","HOLD / NO TRADE"}
