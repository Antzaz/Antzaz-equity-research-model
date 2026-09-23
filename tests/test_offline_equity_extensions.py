from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from offline_equity_extensions import ensure_offline_equity_extensions


def _build_workbook():
    wb=Workbook()
    wb.remove(wb.active)
    d=wb.create_sheet("Company Data")
    d["B4"]="TEST"; d["B8"]=100.0; d["B9"]=1.0; d["B10"]=100.0
    d["B12"]=20.0; d["B13"]=5.0; d["B15"]=20.0

    h=wb.create_sheet("Historical Financials")
    for c,(year,rev,op,ni,eps,ocf,capex,depr,sbc) in enumerate([
        (2024,100.0,20.0,15.0,5.0,24.0,8.0,4.0,2.0),
        (2025,110.0,24.0,18.0,6.0,28.0,9.0,4.5,2.2),
    ],2):
        h.cell(3,c,year); h.cell(4,c,rev); h.cell(9,c,op); h.cell(11,c,ni)
        h.cell(12,c,eps); h.cell(14,c,ocf); h.cell(15,c,capex)
        h.cell(18,c,depr); h.cell(21,c,sbc)

    e=wb.create_sheet("Expectations & Consensus")
    rows=[
        ("Revenue",2026,120.0,125.0),
        ("EPS",2026,6.5,6.8),
        ("EBIT Margin",2026,0.225,0.23),
        ("FCF",2026,21.0,22.0),
    ]
    for r,row in enumerate(rows,7):
        e.cell(r,1,row[0]); e.cell(r,2,row[1]); e.cell(r,3,row[2]); e.cell(r,4,row[3])
        e.cell(r,7,0.01); e.cell(r,8,0.02); e.cell(r,9,0.05); e.cell(r,12,2); e.cell(r,13,"Provider A + B")

    a=wb.create_sheet("Advanced Analytics")
    a["A7"]=2024; a["B7"]=80.0; a["C7"]=5.0; a["D7"]=16.0
    a["A8"]=2025; a["B8"]=95.0; a["C8"]=6.0; a["D8"]=95.0/6.0
    a["I7"]="2026-01-20"; a["J7"]=1.50; a["K7"]=1.60; a["L7"]=0.0667

    c=wb.create_sheet("Cost of Capital")
    c["A1"]="Base WACC"; c["B1"]=0.09
    c["A2"]="Effective tax rate"; c["B2"]=0.20

    fs=wb.create_sheet("Financial Statements")
    fs["A1"]="Metric"; fs["B1"]=2024; fs["C1"]=2025
    data=[
        ("Total Debt",10.0,9.0),
        ("Total Equity",60.0,68.0),
        ("Cash and Cash Equivalents",15.0,20.0),
        ("Repurchase Of Capital Stock",3.0,4.0),
        ("Dividends Paid",1.0,1.2),
        ("Business Acquisitions",2.0,1.0),
        ("Repayments of Debt",1.0,1.5),
    ]
    for r,row in enumerate(data,2):
        fs.cell(r,1,row[0]); fs.cell(r,2,row[1]); fs.cell(r,3,row[2])

    s=wb.create_sheet("Segment Analysis")
    s["A1"]="Segment"; s["B1"]=2024; s["C1"]=2025
    s["A2"]="Core"; s["B2"]=70.0; s["C2"]=77.0
    s["A3"]="Growth"; s["B3"]=30.0; s["C3"]=33.0

    wb.create_sheet("Data Quality")
    return wb


def test_offline_equity_extensions_create_private_research_sheets_and_history(tmp_path):
    root=tmp_path/"research_data"
    wb=_build_workbook()
    first=ensure_offline_equity_extensions(
        wb,"TEST",research_root=root,persist_history=True,captured_at="2026-01-01T00:00:00+00:00"
    )
    required={
        "Forecast Accountability","Earnings & Revisions","Capital Allocation",
        "Valuation History","SOTP Framework","Thesis Timeline",
    }
    assert required.issubset(set(wb.sheetnames))
    assert first["forecast_snapshots"]==1
    assert first["sotp"]["segments"]==2
    assert (root/"TEST"/"forecast_history.json").exists()
    assert (root/"TEST"/"forecast_accuracy_summary.csv").exists()

    # A second run with revised expectations must create a new revision snapshot.
    wb["Expectations & Consensus"]["C7"]=124.0
    second=ensure_offline_equity_extensions(
        wb,"TEST",research_root=root,persist_history=True,captured_at="2026-06-01T00:00:00+00:00"
    )
    assert second["forecast_snapshots"]==2

    # Once FY2026 actuals are present, prior 2026 forecasts mature automatically.
    h=wb["Historical Financials"]; c=4
    h.cell(3,c,2026); h.cell(4,c,123.0); h.cell(9,c,28.0); h.cell(11,c,20.0)
    h.cell(12,c,6.7); h.cell(14,c,31.0); h.cell(15,c,9.5); h.cell(18,c,5.0); h.cell(21,c,2.4)
    third=ensure_offline_equity_extensions(
        wb,"TEST",research_root=root,persist_history=True,captured_at="2027-02-01T00:00:00+00:00"
    )
    assert third["matured_forecast_records"]>=8
    assert wb["Forecast Accountability"].max_row>10
    assert wb["Capital Allocation"]["B18"].value is not None or wb["Capital Allocation"].max_row>=18
    assert wb["Data Quality"].max_row>=4


def test_sotp_framework_never_invents_segment_multiple(tmp_path):
    wb=_build_workbook()
    ensure_offline_equity_extensions(
        wb,"TEST",research_root=tmp_path/"r",persist_history=False,captured_at="2026-01-01T00:00:00+00:00"
    )
    ws=wb["SOTP Framework"]
    # Multiples are yellow analyst-input cells and intentionally blank.
    assert ws["D7"].value is None
    assert ws["D8"].value is None
    assert "REVIEW" in str(ws["G7"].value)


def test_first_forecast_run_does_not_claim_accuracy(tmp_path):
    wb=_build_workbook()
    out=ensure_offline_equity_extensions(
        wb,"TEST",research_root=tmp_path/"r",persist_history=True,captured_at="2026-01-01T00:00:00+00:00"
    )
    assert out["matured_forecast_records"]==0
    ws=wb["Forecast Accountability"]
    assert any("No matured stored forecast" in str(ws.cell(r,2).value or "") for r in range(1,ws.max_row+1))


def test_sotp_uses_sector_appropriate_direct_value_mode_for_bank(tmp_path):
    wb=_build_workbook()
    wb["Company Data"]["B6"]="Financial Services"
    wb["Company Data"]["B7"]="Banks - Diversified"
    out=ensure_offline_equity_extensions(
        wb,"TESTBANK",research_root=tmp_path/"r",persist_history=False,
        captured_at="2026-01-01T00:00:00+00:00"
    )
    assert out["sotp"]["policy"]=="bank"
    assert out["sotp"]["direct_value_mode"] is True
    ws=wb["SOTP Framework"]
    assert ws["D6"].value=="Analyst Segment Value"
    assert ws["E7"].value=='=IFERROR(D7,"")'


def test_same_day_forecast_refresh_replaces_snapshot(tmp_path):
    from offline_equity_extensions import _load_history

    root=tmp_path/"research_data"
    wb=_build_workbook()
    ensure_offline_equity_extensions(
        wb,"TEST",research_root=root,persist_history=True,
        captured_at="2026-01-01T08:00:00+00:00"
    )
    wb["Expectations & Consensus"]["C7"]=130.0
    ensure_offline_equity_extensions(
        wb,"TEST",research_root=root,persist_history=True,
        captured_at="2026-01-01T16:00:00+00:00"
    )
    history=_load_history(root,"TEST")
    assert len(history["snapshots"])==1
    revenue=next(x for x in history["snapshots"][0]["estimates"] if x["Metric"]=="Revenue")
    assert revenue["Consensus"]==130.0
