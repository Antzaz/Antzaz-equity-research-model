from __future__ import annotations

import finland_sources
import issuer_source_engine
import issuer_statement_profiles
import consensus_sources
import wacc_engine
import dynamic_peer_engine
from source_registry import full_source_catalog, market_sources


def _sample_esef_payload():
    return {
        "documentInfo": {"documentType": "https://xbrl.org/2021/xbrl-json"},
        "facts": {
            "rev": {
                "value": "1000000000",
                "dimensions": {
                    "concept": "ifrs-full:Revenue",
                    "entity": "lei:TEST",
                    "period": "2024-01-01T00:00:00/2025-01-01T00:00:00",
                    "unit": "iso4217:EUR",
                },
            },
            "op": {
                "value": "140000000",
                "dimensions": {
                    "concept": "ifrs-full:ProfitLossFromOperatingActivities",
                    "entity": "lei:TEST",
                    "period": "2024-01-01T00:00:00/2025-01-01T00:00:00",
                    "unit": "iso4217:EUR",
                },
            },
            "ni": {
                "value": "100000000",
                "dimensions": {
                    "concept": "ifrs-full:ProfitLoss",
                    "entity": "lei:TEST",
                    "period": "2024-01-01T00:00:00/2025-01-01T00:00:00",
                    "unit": "iso4217:EUR",
                },
            },
            "eps": {
                "value": "2.0",
                "dimensions": {
                    "concept": "ifrs-full:DilutedEarningsLossPerShare",
                    "entity": "lei:TEST",
                    "period": "2024-01-01T00:00:00/2025-01-01T00:00:00",
                    "unit": "iso4217:EUR/xbrli:shares",
                },
            },
            "assets": {
                "value": "2500000000",
                "dimensions": {
                    "concept": "ifrs-full:Assets",
                    "entity": "lei:TEST",
                    "period": "2025-01-01T00:00:00",
                    "unit": "iso4217:EUR",
                },
            },
            # A segment/dimensioned fact must never overwrite consolidated history.
            "segment_rev": {
                "value": "9000000000",
                "dimensions": {
                    "concept": "ifrs-full:Revenue",
                    "entity": "lei:TEST",
                    "period": "2024-01-01T00:00:00/2025-01-01T00:00:00",
                    "unit": "iso4217:EUR",
                    "test:SegmentAxis": "test:SegmentMember",
                },
            },
        },
    }


def test_finland_source_routing_is_specific_to_helsinki():
    assert finland_sources.is_finnish_ticker("KNEBV.HE")
    assert finland_sources.is_finnish_ticker("nokia.he")
    assert not finland_sources.is_finnish_ticker("NOK")
    sources=market_sources("KNEBV.HE")
    assert "finland_oam" in sources
    assert "xbrl_international_esef" in sources
    assert "statistics_finland" in sources
    assert market_sources("MSFT")=={}
    catalog=full_source_catalog("KNEBV.HE")
    assert catalog["market"]["bank_of_finland"]["env_key"] is None


def test_esef_parser_keeps_consolidated_annual_facts_only():
    hist,meta=finland_sources._extract_esef_report(_sample_esef_payload())
    assert hist[2024]["revenue"]==1_000_000_000
    assert hist[2024]["op"]==140_000_000
    assert hist[2024]["ni"]==100_000_000
    assert hist[2024]["eps"]==2.0
    assert hist[2024]["shares"]==50_000_000
    assert hist[2024]["assets"]==2_500_000_000
    assert meta["years"]==[2024]


def test_finland_esef_history_uses_free_json_filing(monkeypatch):
    monkeypatch.setattr(
        finland_sources,
        "_filing_json_urls",
        lambda ticker,info,limit=8: (
            ["https://filings.xbrl.org/test.json"],
            {"status":"PASS","matched_name":"test"},
        ),
    )
    monkeypatch.setattr(
        finland_sources,
        "_request_json",
        lambda url,params=None,timeout=20: _sample_esef_payload(),
    )
    hist,meta=finland_sources.finland_esef_history(
        "TEST.HE",
        {"longName":"Test Oyj"},
    )
    assert hist[2024]["revenue"]==1_000_000_000
    assert meta["status"]=="PASS"
    assert "Finnish OAM" in meta["source"]


def test_crossborder_history_prioritizes_finnish_esef(monkeypatch):
    monkeypatch.setattr(
        issuer_source_engine,
        "finland_esef_history",
        lambda ticker,info: (
            {2024:{"revenue":100.0,"ni":10.0}},
            {"status":"PASS","years":[2024]},
        ),
    )
    monkeypatch.setattr(
        issuer_source_engine,
        "issuer_website_history",
        lambda ticker,info: (
            {2024:{"revenue":90.0,"op":12.0}},
            {"source":"issuer"},
        ),
    )
    monkeypatch.setattr(
        issuer_source_engine,
        "sec_crossborder_history",
        lambda facts: ({},{"source":"sec"}),
    )
    monkeypatch.setattr(
        issuer_source_engine,
        "yfinance_statement_history",
        lambda ticker: (
            {2024:{"revenue":80.0,"op":11.0,"cash":5.0}},
            {"source":"yahoo"},
        ),
    )

    hist,meta=issuer_source_engine.build_crossborder_history(
        "TEST.HE",
        {"longName":"Test Oyj"},
        None,
    )
    assert hist[2024]["revenue"]==100.0
    assert hist[2024]["op"]==12.0
    assert hist[2024]["cash"]==5.0
    assert meta["priority"][0].startswith("Finnish OAM")


def test_guarded_filings_wrapper_forwards_company_info(monkeypatch):
    import safe_update_model
    from openpyxl import Workbook

    captured={}
    def fake_update_filings(wb,ticker,info=None):
        captured["ticker"]=ticker
        captured["info"]=info
        return "ok"

    monkeypatch.setattr(safe_update_model,"_ORIGINAL_UPDATE_FILINGS",fake_update_filings)
    wb=Workbook()
    wb.active.title="Filings"
    result=safe_update_model._safe_update_filings(
        wb,
        "YIT.HE",
        {"longName":"YIT Oyj"},
    )
    assert result=="ok"
    assert captured["ticker"]=="YIT.HE"
    assert captured["info"]["longName"]=="YIT Oyj"


def test_update_model_compat_accepts_legacy_two_arg_wrapper(monkeypatch):
    import update_model
    from openpyxl import Workbook

    called={}
    def legacy_wrapper(wb,ticker):
        called["ticker"]=ticker
        return "legacy-ok"

    monkeypatch.setattr(update_model,"update_filings",legacy_wrapper)
    wb=Workbook()
    wb.active.title="Filings"
    result=update_model._call_update_filings_compat(
        wb,
        "YIT.HE",
        {"longName":"YIT Oyj"},
    )
    assert result=="legacy-ok"
    assert called["ticker"]=="YIT.HE"


def test_update_model_compat_forwards_info_to_new_wrapper(monkeypatch):
    import update_model
    from openpyxl import Workbook

    called={}
    def new_wrapper(wb,ticker,info=None):
        called["ticker"]=ticker
        called["info"]=info
        return "new-ok"

    monkeypatch.setattr(update_model,"update_filings",new_wrapper)
    wb=Workbook()
    wb.active.title="Filings"
    result=update_model._call_update_filings_compat(
        wb,
        "YIT.HE",
        {"longName":"YIT Oyj"},
    )
    assert result=="new-ok"
    assert called["info"]["longName"]=="YIT Oyj"


def test_esef_companyfacts_adapter_exposes_standard_ifrs_facts():
    facts=finland_sources._esef_companyfacts_from_payload(_sample_esef_payload(),"https://example.test/yit.json")
    assert facts["_source_type"]=="Finnish OAM / ESEF"
    revenue=facts["facts"]["ifrs-full"]["Revenue"]["units"]["EUR"]
    assert revenue[0]["form"]=="ESEF"
    assert revenue[0]["fy"]==2024
    assert revenue[0]["val"]==1_000_000_000
    assert facts["facts"]["ifrs-full"]["DilutedEarningsLossPerShare"]["units"]["EUR/shares"][0]["val"]==2.0


def test_dimensioned_esef_segment_parser_keeps_segment_economics():
    payload={
        "facts":{
            "seg_rev":{
                "value":"317000000",
                "dimensions":{
                    "concept":"ifrs-full:Revenue",
                    "entity":"lei:TEST",
                    "period":"2025-01-01T00:00:00/2026-01-01T00:00:00",
                    "unit":"iso4217:EUR",
                    "ifrs-full:SegmentsAxis":"yit:ResidentialFinlandMember",
                },
            },
            "seg_profit":{
                "value":"-8000000",
                "dimensions":{
                    "concept":"yit:AdjustedOperatingProfit",
                    "entity":"lei:TEST",
                    "period":"2025-01-01T00:00:00/2026-01-01T00:00:00",
                    "unit":"iso4217:EUR",
                    "ifrs-full:SegmentsAxis":"yit:ResidentialFinlandMember",
                },
            },
        }
    }
    segments=finland_sources._extract_esef_segments(payload)
    rec=segments["Residential Finland"]
    assert rec["revenue"][2025]==317_000_000
    assert rec["op"][2025]==-8_000_000


def test_helsinki_tickers_use_ifrs_statement_profile():
    profile=issuer_statement_profiles.get_statement_profile("YIT.HE")
    assert profile["key"]=="ifrs"
    assert profile["name"]=="IFRS / ESEF corporate"
    assert any(row[0]=="Operating Income" for row in profile["income"])
    assert any(row[0]=="Cash & Cash Equivalents" for row in profile["balance"])


def test_finnish_core_extension_fact_maps_to_statutory_ifrs_row():
    payload={
        "facts":{
            "op":{
                "value":"45000000",
                "dimensions":{
                    "concept":"yit:OperatingProfit",
                    "entity":"lei:TEST",
                    "period":"2025-01-01T00:00:00/2026-01-01T00:00:00",
                    "unit":"iso4217:EUR",
                },
            },
            "adjusted":{
                "value":"54000000",
                "dimensions":{
                    "concept":"yit:AdjustedOperatingProfit",
                    "entity":"lei:TEST",
                    "period":"2025-01-01T00:00:00/2026-01-01T00:00:00",
                    "unit":"iso4217:EUR",
                },
            },
        }
    }
    facts=finland_sources._esef_companyfacts_from_payload(payload)
    op=facts["facts"]["ifrs-full"]["ProfitLossFromOperatingActivities"]["units"]["EUR"]
    assert op[0]["val"]==45_000_000
    # Adjusted/non-GAAP operating profit must never silently become statutory operating profit.
    assert all("AdjustedOperatingProfit" not in key for key in facts["facts"]["ifrs-full"])


def test_vara_detailed_consensus_parser_handles_fy_columns():
    import pandas as pd
    df=pd.DataFrame([
        ["Revenue","Number of Estimates",6,6,6],
        ["Revenue","Average",1973.7,2279.8,2529.5],
        ["Revenue","Highest",2040.0,2545.0,2761.0],
        ["Revenue","Lowest",1927.0,2093.0,2395.0],
        ["Adjusted operating profit","Average",87.6,126.2,158.2],
    ],columns=["KPI","[EUR m]","FY 2026 E","FY 2027 E","FY 2028 E"])
    rows=finland_sources._parse_vara_detailed_tables([df])
    keyed={(x["metric"],x["fiscal_year"]):x for x in rows}
    assert keyed[("Revenue",2026)]["mean"]==1973.7
    assert keyed[("Revenue",2026)]["analysts"]==6
    assert keyed[("Revenue",2026)]["high"]==2040.0
    assert keyed[("Adjusted Operating Profit",2027)]["mean"]==126.2


def test_vara_price_target_parser():
    import pandas as pd
    df=pd.DataFrame([
        ["Price, last closing",4.24],
        ["Number of Estimates",6],
        ["Highest",5.00],
        ["Average",4.45],
        ["Median",4.40],
        ["Lowest",4.10],
    ],columns=["Metric","Current"])
    target=finland_sources._parse_vara_price_target_tables([df])
    assert target["mean"]==4.45
    assert target["median"]==4.40
    assert target["analysts"]==6


def test_vara_records_keep_adjusted_profit_separate(monkeypatch):
    monkeypatch.setattr(
        consensus_sources,
        "finland_vara_consensus",
        lambda ticker,info: (
            [
                {"metric":"Revenue","fiscal_year":2026,"mean":1973.7,"low":1927.0,"high":2040.0,"analysts":6},
                {"metric":"Adjusted Operating Profit","fiscal_year":2026,"mean":87.6,"low":83.0,"high":96.0,"analysts":6},
            ],
            {"mean":4.45,"median":4.40,"low":4.10,"high":5.0,"analysts":6,"source_url":"https://vara.example/yit"},
            {"status":"PASS","source_url":"https://vara.example/yit"},
        ),
    )
    records,status,target=consensus_sources.vara_records(
        "YIT.HE",2025,{"currency":"EUR","financialCurrency":"EUR"}
    )
    metrics={r["metric"]:r for r in records}
    assert abs(metrics["Revenue"]["mean"]-1.9737)<1e-9
    assert abs(metrics["Adjusted Operating Profit"]["mean"]-0.0876)<1e-9
    assert "EBIT" not in metrics
    assert status["ok"]
    assert target["mean"]==4.45


def test_helsinki_wacc_routes_to_eur_risk_free(monkeypatch):
    monkeypatch.setattr(
        wacc_engine,
        "_ecb_eur_10y",
        lambda: (0.031,"2026-10-20","ECB AAA euro-area 10Y spot rate"),
    )
    value,asof,status,source,role=wacc_engine._risk_free(
        "YIT.HE",{"currency":"EUR","financialCurrency":"EUR"}
    )
    assert value==0.031
    assert "ECB" in status
    assert "EUR" in role
    assert "data.ecb.europa.eu" in source


def test_helsinki_wacc_uses_finland_tax_fallback(monkeypatch):
    from openpyxl import Workbook
    wb=Workbook()
    wb.active.title="Company Data"
    wb["Company Data"]["B10"]=1.0
    wb["Company Data"]["B13"]=0.0
    monkeypatch.setattr(wacc_engine,"_risk_free",lambda ticker,info:(0.03,"2026-10-20","ECB","https://ecb.example","EUR"))
    monkeypatch.setattr(wacc_engine,"_damodaran_erp",lambda:(0.04,"2026-10-20","ERP"))
    monkeypatch.setattr(wacc_engine,"_beta",lambda ticker,info:(1.0,"test"))
    monkeypatch.setattr(wacc_engine,"_credit_spread",lambda wb:(0.01,"test"))
    result=wacc_engine.calculate_wacc(wb,"YIT.HE",{"currency":"EUR"})
    assert result["tax_rate"]==0.20
    assert result["currency"]=="EUR"


def test_nordic_construction_peers_are_prioritized_for_yit():
    peers=dynamic_peer_engine.STRATEGIC_PEERS["YIT.HE"]
    assert peers[:4]==["SKA-B.ST","NCC-B.ST","PEAB-B.ST","SRV1V.HE"]
    assert "CONSTI.HE" in peers
    assert "KREATE.HE" in peers
    assert "GRK.HE" in peers
    assert dynamic_peer_engine.NORDIC_INDUSTRY_FALLBACK["Engineering & Construction"][0]=="SKA-B.ST"


def test_finland_sector_fallbacks_precede_us_fallback_candidates(monkeypatch):
    monkeypatch.setattr(dynamic_peer_engine,"preferred_peer_symbols",lambda ticker,industry:[])
    monkeypatch.setattr(dynamic_peer_engine,"yf",None)
    out=dynamic_peer_engine._discover_candidates(
        "YIT.HE","Industrials","Engineering & Construction","","", "FI"
    )
    assert out.index("SKA-B.ST") < out.index("ETN")
    assert out.index("SRV1V.HE") < out.index("ETN")
