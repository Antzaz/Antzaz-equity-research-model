from pathlib import Path

import pandas as pd
from openpyxl import Workbook

import research_accountability as ra


def _workbook():
    wb=Workbook()
    ws=wb.active; ws.title="Company Data"
    ws["B4"]="TEST"; ws["B8"]=100.0; ws["B9"]=1.0; ws["B14"]=-5.0

    h=wb.create_sheet("Historical Financials")
    years=[2020,2021,2022,2023,2024,2025]
    revenue=[100,112,125,140,158,180]
    op=[20,23,26,31,37,45]
    ni=[15,17,19,23,29,34]
    eps=[15,17,19,23,29,34]
    ocf=[22,25,29,34,42,50]
    capex=[5,6,7,8,9,10]
    sbc=[2,2.2,2.5,2.8,3.0,3.2]
    for i,y in enumerate(years,2):
        h.cell(3,i,y); h.cell(4,i,revenue[i-2]); h.cell(9,i,op[i-2]); h.cell(11,i,ni[i-2])
        h.cell(12,i,eps[i-2]); h.cell(14,i,ocf[i-2]); h.cell(15,i,capex[i-2]); h.cell(21,i,sbc[i-2])

    s=wb.create_sheet("Three-Case Scenarios")
    s["C6"]=0.09
    for c in range(14,24):
        s.cell(12,c,0.08-(c-14)*0.003)
        s.cell(14,c,0.25)
        s.cell(18,c,0.03)
        s.cell(20,c,0.05)

    fs=wb.create_sheet("Financial Statements")
    fs["A2"]="Metric"
    for i,y in enumerate(years,2): fs.cell(2,i,y)
    rows={
        3:("Total Assets",[130,145,160,178,198,220]),
        4:("Cash and Cash Equivalents",[20,22,25,28,32,36]),
        5:("Total Current Liabilities",[30,32,35,38,41,45]),
        6:("Repurchase of Capital Stock",[2,3,4,5,6,7]),
        7:("Dividends Paid",[1,1.2,1.4,1.6,1.8,2.0]),
        8:("Payments to Acquire Businesses, Net of Cash Acquired",[0.5,0.4,0.8,0.6,0.7,0.9]),
    }
    for r,(label,vals) in rows.items():
        fs.cell(r,1,label)
        for i,v in enumerate(vals,2): fs.cell(r,i,v)

    seg=wb.create_sheet("Segment Analysis")
    seg.append(["Segment","2023","2024","2025"])
    seg.append(["Core",50,55,60])
    seg.append(["Cloud",20,28,38])
    seg.append(["Other",5,6,7])

    con=wb.create_sheet("Expectations & Consensus")
    con["A1"]="Metric"; con["B1"]="Current"
    con["A2"]="Revenue Growth"; con["B2"]=0.10
    con["A3"]="Forward EPS"; con["B3"]=5.0

    dq=wb.create_sheet("Data Quality")
    dq.append(["Control","Status","Detail"])

    dv=wb.create_sheet("Decision View")
    dv["A1"]="Current Market Price"; dv["B1"]=100.0
    dv["A2"]="Base DCF Fair Value"; dv["B2"]=125.0
    dv["A3"]="MODEL VIEW"; dv["B3"]="POTENTIALLY ATTRACTIVE"
    return wb


def test_research_accountability_builds_all_layers(tmp_path, monkeypatch):
    wb=_workbook()
    monkeypatch.setattr(ra,"RESEARCH_DATA",tmp_path/"research_data")
    dates=pd.to_datetime(["2020-12-31","2021-12-31","2022-12-30","2023-12-29","2024-12-31","2025-12-31"])
    monkeypatch.setattr(ra,"_price_history",lambda ticker: pd.Series([40,48,55,67,82,100],index=dates))
    monkeypatch.setattr(ra,"_earnings_history",lambda ticker: pd.DataFrame([{"quarter":"2025Q4","epsEstimate":1.0,"epsActual":1.1,"surprisePercent":10.0}]))

    result=ra.apply_research_accountability(
        wb,"TEST",
        {"forwardEps":5.0,"forwardPE":20.0,"targetMeanPrice":130.0,"revenueGrowth":0.1,"earningsGrowth":0.12,"effectiveTaxRate":0.20},
    )

    expected={
        "Forecast Accountability","Earnings & Revisions","Capital Allocation",
        "Valuation History","SOTP Framework","Thesis Timeline",
    }
    assert expected.issubset(set(wb.sheetnames))
    assert (tmp_path/"research_data"/"TEST"/"forecast_history.csv").exists()
    assert (tmp_path/"research_data"/"TEST"/"estimate_history.csv").exists()
    assert (tmp_path/"research_data"/"TEST"/"sotp_assumptions.csv").exists()
    assert (tmp_path/"research_data"/"TEST"/"thesis_timeline.csv").exists()
    assert result["capital_allocation"]["history"].shape[0] == 6
    assert len(result["valuation_history"]["history"]) >= 3


def test_forecasts_only_score_when_actual_exists(tmp_path, monkeypatch):
    wb=_workbook()
    monkeypatch.setattr(ra,"RESEARCH_DATA",tmp_path/"research_data")
    cur=ra._forecast_rows(wb)
    hist=ra._append_forecast_snapshot("TEST",cur)
    detail,summary=ra._forecast_accuracy(hist,ra._historical(wb))
    # Forecasts start after the latest actual year, so there must be no fake accuracy yet.
    assert detail.empty
    assert summary.empty


def test_sotp_requires_explicit_multiples(tmp_path, monkeypatch):
    wb=_workbook()
    monkeypatch.setattr(ra,"RESEARCH_DATA",tmp_path/"research_data")
    first=ra.ensure_sotp_framework(wb,"TEST")
    assert first["valid"] == 0
    path=first["path"]
    df=pd.read_csv(path)
    df["Multiple"]=[5.0+i for i in range(len(df))]
    df.to_csv(path,index=False)
    second=ra.ensure_sotp_framework(wb,"TEST")
    assert second["valid"] >= 2
    assert second["value_per_share"] is not None


def test_display_path_is_non_recursive(tmp_path: Path):
    inside=ra.BASE / "research_data" / "TEST" / "sample.csv"
    shown=ra._display_path(inside)
    assert "research_data" in shown
    assert shown.endswith("sample.csv")
    outside=tmp_path / "x.csv"
    assert ra._display_path(outside).endswith("x.csv")
