from __future__ import annotations

"""Private/offline research-accountability extensions.

This module closes the loop between a model build and later evidence. It is conservative:
- forecasts are snapshotted point-in-time and compared with actuals only when actuals exist;
- estimate revisions are stored locally across builds;
- capital-allocation metrics use disclosed statement rows and label approximations explicitly;
- valuation-history statistics are only shown when usable price/fundamental history exists;
- SOTP stays REVIEW until segment economics and explicit local multiples exist;
- thesis events are durable local research records, not generated narrative.

The module never changes core valuation inputs or portfolio positions.
"""

from datetime import datetime
from pathlib import Path
import math
import re

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

try:
    import yfinance as yf
except Exception:
    yf = None


BASE = Path(__file__).resolve().parent
RESEARCH_DATA = BASE / "research_data"

NAVY = "17365D"
BLUE = "2F75B5"
WHITE = "FFFFFF"
GREEN = "E2F0D9"
GOLD = "FFF2CC"
RED = "FCE4D6"
GREY = "666666"
LIGHT = "F5F9FC"
FMT_PCT = '0.0%;[Red](0.0%);-'
FMT_PRICE = '$#,##0.00;[Red]($#,##0.00);-'
FMT_BN = '#,##0.0;[Red](#,##0.0);-'
FMT_MULT = '0.0x;[Red](0.0x);-'
FMT_SCORE = '0.0'


def _fill(color: str):
    return PatternFill("solid", fgColor=color)


def _num(value, default=None):
    try:
        if isinstance(value, bool) or value in (None, ""):
            return default
        x = float(value)
        return x if math.isfinite(x) else default
    except Exception:
        return default


def _text(value):
    if value in (None, ""):
        return ""
    return str(value).strip()


def _find(ws, labels, col=1):
    if isinstance(labels, str):
        labels = [labels]
    needles = {str(x).strip().lower() for x in labels}
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, col).value or "").strip().lower() in needles:
            return r
    return None


def _sheet(wb, name):
    if name in wb.sheetnames:
        wb.remove(wb[name])
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    return ws


def _title(ws, title, subtitle):
    for c in range(1, 9):
        ws.cell(1, c).fill = _fill(NAVY)
        ws.cell(2, c).fill = _fill(NAVY)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, color=WHITE, size=18)
    ws["A3"] = subtitle
    ws.merge_cells("A3:H3")
    ws["A3"].font = Font(italic=True, color=GREY)
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")


def _section(ws, row, title, end=8):
    for c in range(1, end + 1):
        ws.cell(row, c).fill = _fill(NAVY)
        ws.cell(row, c).font = Font(bold=True, color=WHITE)
    ws.cell(row, 1, title)


def _header(ws, row, values):
    for c, value in enumerate(values, 1):
        ws.cell(row, c, value)
        ws.cell(row, c).fill = _fill(BLUE)
        ws.cell(row, c).font = Font(bold=True, color=WHITE)
        ws.cell(row, c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _status_fill(status):
    return _fill(GREEN if status == "PASS" else RED if status == "FAIL" else GOLD)


def _quality_row(wb, label, status, detail):
    if "Data Quality" not in wb.sheetnames:
        return
    ws = wb["Data Quality"]
    r = _find(ws, label)
    if not r:
        r = ws.max_row + 1
    ws.cell(r, 1, label)
    ws.cell(r, 2, status)
    ws.cell(r, 2).fill = _status_fill(status)
    ws.cell(r, 2).font = Font(bold=True)
    ws.cell(r, 3, detail)
    ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")


def _ticker_dir(ticker):
    path = RESEARCH_DATA / str(ticker).upper()
    path.mkdir(parents=True, exist_ok=True)
    return path


def _display_path(path: Path) -> str:
    """Return a stable human-readable private path without recursing or exposing more than needed."""
    p=Path(path)
    try:
        return str(p.resolve().relative_to(BASE.resolve()))
    except Exception:
        try:
            return str(p.relative_to(BASE))
        except Exception:
            return str(p)


def _historical(wb):
    if "Historical Financials" not in wb.sheetnames:
        return pd.DataFrame()
    ws = wb["Historical Financials"]
    rows = {
        "Revenue": 4,
        "OperatingIncome": 9,
        "NetIncome": 11,
        "EPS": 12,
        "OperatingCashFlow": 14,
        "Capex": 15,
        "SBC": 21,
    }
    out = []
    for c in range(2, 8):
        y = _num(ws.cell(3, c).value)
        if not y or not (1900 <= int(y) <= 2100):
            continue
        row = {"Year": int(y)}
        for key, rr in rows.items():
            row[key] = _num(ws.cell(rr, c).value)
        if row["OperatingCashFlow"] is not None and row["Capex"] is not None:
            row["FCF"] = row["OperatingCashFlow"] - abs(row["Capex"])
        else:
            row["FCF"] = None
        if row["NetIncome"] is not None and row["EPS"] not in (None, 0):
            row["DilutedShares"] = row["NetIncome"] / row["EPS"]
        else:
            row["DilutedShares"] = None
        out.append(row)
    return pd.DataFrame(out).sort_values("Year") if out else pd.DataFrame()


def _current_price(wb):
    if "Company Data" in wb.sheetnames:
        return _num(wb["Company Data"]["B8"].value)
    return None


def _shares(wb):
    if "Company Data" in wb.sheetnames:
        return _num(wb["Company Data"]["B9"].value)
    return None


def _net_debt(wb):
    if "Company Data" in wb.sheetnames:
        return _num(wb["Company Data"]["B14"].value)
    return None


def _decision_values(wb):
    out = {"CurrentPrice": _current_price(wb), "BaseValue": None, "BearValue": None, "BullValue": None, "ModelView": None}
    if "Decision View" in wb.sheetnames:
        ws = wb["Decision View"]
        for labels, key in [
            (["Current Market Price"], "CurrentPrice"),
            (["Base DCF Fair Value"], "BaseValue"),
            (["MODEL VIEW"], "ModelView"),
        ]:
            r = _find(ws, labels)
            if r:
                out[key] = ws.cell(r, 2).value
        # Severe bear is useful when a conventional bear value is unavailable.
        r = _find(ws, ["Monte Carlo P10", "Severe Bear / Share"])
        if r:
            out["BearValue"] = ws.cell(r, 2).value
    if "Investment Summary" in wb.sheetnames:
        ws = wb["Investment Summary"]
        mapping = {
            "CurrentPrice": ["Current Price"],
            "BaseValue": ["Base DCF / Share"],
            "BearValue": ["Severe Bear / Share", "Bear Value / Share"],
            "BullValue": ["Bull Value / Share"],
            "ModelView": ["Model View"],
        }
        for key, labels in mapping.items():
            if out.get(key) not in (None, ""):
                continue
            r = _find(ws, labels)
            if r:
                out[key] = ws.cell(r, 2).value
    for key in ("CurrentPrice", "BaseValue", "BearValue", "BullValue"):
        out[key] = _num(out.get(key))
    out["ModelView"] = _text(out.get("ModelView")) or None
    return out


def _base_wacc(wb):
    if "Cost of Capital" in wb.sheetnames:
        ws = wb["Cost of Capital"]
        # Prefer the scenario table's Base row.
        r = _find(ws, ["Base"])
        if r:
            for c in range(2, min(6, ws.max_column) + 1):
                v = _num(ws.cell(r, c).value)
                if v is not None and 0 < v < 1:
                    return v
    if "Three-Case Scenarios" in wb.sheetnames:
        return _num(wb["Three-Case Scenarios"]["C6"].value)
    return None


def _statement_series(wb, labels):
    if "Financial Statements" not in wb.sheetnames:
        return {}
    ws = wb["Financial Statements"]
    row = None
    for label in labels:
        row = _find(ws, label)
        if row:
            break
    if not row:
        return {}
    # Search upward for the closest Metric/year header.
    header = None
    for rr in range(row - 1, max(0, row - 80), -1):
        if str(ws.cell(rr, 1).value or "").strip().lower() == "metric":
            header = rr
            break
    if not header:
        return {}
    out = {}
    for c in range(2, min(ws.max_column, 14) + 1):
        y = _num(ws.cell(header, c).value)
        v = _num(ws.cell(row, c).value)
        if y and v is not None and 1900 <= int(y) <= 2100:
            out[int(y)] = v
    return out


def _forecast_rows(wb):
    hist = _historical(wb)
    if hist.empty or "Three-Case Scenarios" not in wb.sheetnames:
        return pd.DataFrame()
    ws = wb["Three-Case Scenarios"]
    latest = hist.iloc[-1]
    latest_year = int(latest["Year"])
    revenue = _num(latest["Revenue"])
    op_income = _num(latest["OperatingIncome"])
    fcf = _num(latest["FCF"])
    if revenue is None:
        return pd.DataFrame()
    cash_conversion = None
    if op_income not in (None, 0) and fcf is not None:
        cash_conversion = max(0.15, min(1.35, fcf / op_income))
    else:
        cash_conversion = 0.70

    rows = []
    rev = revenue
    for i, c in enumerate(range(14, 24), 1):  # Base-case block.
        growth = _num(ws.cell(12, c).value)
        margin = _num(ws.cell(14, c).value)
        if growth is None:
            continue
        rev *= 1 + growth
        op = rev * margin if margin is not None else None
        est_fcf = op * cash_conversion if op is not None else None
        year = latest_year + i
        rows.extend([
            {"ForecastYear": year, "Metric": "Revenue", "Forecast": rev},
            {"ForecastYear": year, "Metric": "Operating Margin", "Forecast": margin},
            {"ForecastYear": year, "Metric": "FCF", "Forecast": est_fcf},
        ])
    return pd.DataFrame(rows)


def _append_forecast_snapshot(ticker, current):
    path = _ticker_dir(ticker) / "forecast_history.csv"
    stamp = datetime.now().astimezone().isoformat(timespec="seconds")
    current = current.copy()
    current["BuildTimestamp"] = stamp
    current["BuildDate"] = stamp[:10]
    if path.exists():
        old = pd.read_csv(path)
        combined = pd.concat([old, current], ignore_index=True)
    else:
        combined = current
    combined = combined.drop_duplicates(["BuildDate", "ForecastYear", "Metric"], keep="last")
    combined.to_csv(path, index=False)
    return combined


def _forecast_accuracy(history, actuals):
    if history.empty or actuals.empty:
        return pd.DataFrame(), pd.DataFrame()
    actual_map = {}
    for _, row in actuals.iterrows():
        year = int(row["Year"])
        actual_map[(year, "Revenue")] = _num(row.get("Revenue"))
        revenue = _num(row.get("Revenue"))
        op = _num(row.get("OperatingIncome"))
        actual_map[(year, "Operating Margin")] = op / revenue if revenue not in (None, 0) and op is not None else None
        actual_map[(year, "FCF")] = _num(row.get("FCF"))
    rows = []
    for _, row in history.iterrows():
        key = (int(row["ForecastYear"]), str(row["Metric"]))
        actual = actual_map.get(key)
        forecast = _num(row.get("Forecast"))
        if actual is None or forecast is None:
            continue
        if str(row["Metric"]) == "Operating Margin":
            err = actual - forecast
            pct_err = err
            unit = "percentage points"
        else:
            err = actual - forecast
            pct_err = err / abs(forecast) if abs(forecast) > 1e-12 else np.nan
            unit = "percent"
        rows.append({
            "BuildTimestamp": row.get("BuildTimestamp"),
            "ForecastYear": key[0],
            "Metric": key[1],
            "Forecast": forecast,
            "Actual": actual,
            "Error": err,
            "ErrorPct": pct_err,
            "AbsoluteErrorPct": abs(pct_err) if pd.notna(pct_err) else np.nan,
            "BiasDirection": "Under-forecast" if err > 0 else "Over-forecast" if err < 0 else "Exact",
            "ErrorUnit": unit,
        })
    detail = pd.DataFrame(rows)
    if detail.empty:
        return detail, pd.DataFrame()
    summary = (
        detail.groupby("Metric", as_index=False)
        .agg(
            Observations=("ErrorPct", "count"),
            MeanError=("ErrorPct", "mean"),
            MeanAbsoluteError=("AbsoluteErrorPct", "mean"),
        )
    )
    return detail, summary


def ensure_forecast_accountability(wb, ticker):
    current = _forecast_rows(wb)
    history = _append_forecast_snapshot(ticker, current) if not current.empty else pd.DataFrame()
    actuals = _historical(wb)
    detail, summary = _forecast_accuracy(history, actuals)
    data_dir = _ticker_dir(ticker)
    if not detail.empty:
        detail.to_csv(data_dir / "forecast_accuracy.csv", index=False)
    if not summary.empty:
        summary.to_csv(data_dir / "forecast_accuracy_summary.csv", index=False)

    ws = _sheet(wb, "Forecast Accountability")
    _title(ws, f"{ticker} — Forecast Accountability", "Point-in-time model forecasts are preserved across builds and compared with actual reported results only after those actuals exist.")
    _section(ws, 5, "Current Base-Case Forecast Snapshot")
    _header(ws, 6, ["Forecast Year", "Metric", "Forecast", "Actual (if reported)", "Error", "Status", "Build History", "Notes"])
    latest_build = history["BuildDate"].max() if not history.empty and "BuildDate" in history else None
    cur = history[history["BuildDate"] == latest_build] if latest_build else current
    actual_map = {}
    for _, r in actuals.iterrows():
        y = int(r["Year"])
        actual_map[(y, "Revenue")] = _num(r.get("Revenue"))
        rev = _num(r.get("Revenue")); op = _num(r.get("OperatingIncome"))
        actual_map[(y, "Operating Margin")] = op / rev if rev not in (None, 0) and op is not None else None
        actual_map[(y, "FCF")] = _num(r.get("FCF"))
    for rr, (_, row) in enumerate(cur.head(30).iterrows(), 7):
        y = int(row["ForecastYear"]); metric = str(row["Metric"]); forecast = _num(row["Forecast"]); actual = actual_map.get((y, metric))
        error = actual - forecast if actual is not None and forecast is not None else None
        status = "PASS" if actual is not None else "OPEN"
        values = [y, metric, forecast, actual, error, status, f"{len(history[(history['ForecastYear']==y)&(history['Metric']==metric)])} snapshot(s)" if not history.empty else "Current build", "No actual is backfilled before the fiscal result exists."]
        for c, v in enumerate(values, 1):
            ws.cell(rr, c, v)
        if metric == "Operating Margin":
            for c in (3, 4, 5): ws.cell(rr, c).number_format = FMT_PCT
        else:
            for c in (3, 4, 5): ws.cell(rr, c).number_format = FMT_BN
        ws.cell(rr, 6).fill = _status_fill("PASS" if status == "PASS" else "REVIEW")

    start = max(40, ws.max_row + 3)
    _section(ws, start, "Historical Forecast Accuracy")
    _header(ws, start + 1, ["Metric", "Observations", "Mean Error", "Mean Absolute Error", "Bias", "Interpretation"])
    for rr, (_, row) in enumerate(summary.iterrows(), start + 2):
        bias = "Under-forecast bias" if row["MeanError"] > 0 else "Over-forecast bias" if row["MeanError"] < 0 else "No directional bias"
        values = [row["Metric"], int(row["Observations"]), row["MeanError"], row["MeanAbsoluteError"], bias, "Error statistics use only matured forecasts with reported actuals."]
        for c, v in enumerate(values, 1): ws.cell(rr, c, v)
        ws.cell(rr, 3).number_format = FMT_PCT
        ws.cell(rr, 4).number_format = FMT_PCT

    status = "PASS" if not detail.empty else "REVIEW"
    _quality_row(wb, "Forecast-vs-actual accountability", status, f"{len(history)} point-in-time forecast rows stored; {len(detail)} matured forecast observations evaluated.")
    for col, width in {"A":16, "B":25, "C":18, "D":20, "E":18, "F":13, "G":18, "H":55}.items(): ws.column_dimensions[col].width = width
    ws.freeze_panes = "A7"
    return {"history": history, "detail": detail, "summary": summary}


def _consensus_snapshot(wb, ticker, info):
    rows = []
    stamp = datetime.now().astimezone().isoformat(timespec="seconds")
    simple = {
        "Forward EPS": info.get("forwardEps"),
        "Forward P/E": info.get("forwardPE"),
        "Target Mean Price": info.get("targetMeanPrice"),
        "Revenue Growth": info.get("revenueGrowth"),
        "Earnings Growth": info.get("earningsGrowth"),
    }
    for metric, value in simple.items():
        value = _num(value)
        if value is not None:
            rows.append({"Timestamp": stamp, "Metric": metric, "Period": "Current provider field", "Value": value, "Source": "Yahoo Finance / provider snapshot"})

    if "Expectations & Consensus" in wb.sheetnames:
        ws = wb["Expectations & Consensus"]
        for r in range(1, min(ws.max_row, 120) + 1):
            label = _text(ws.cell(r, 1).value)
            if not label or len(label) > 80:
                continue
            for c in range(2, min(ws.max_column, 8) + 1):
                value = _num(ws.cell(r, c).value)
                if value is None:
                    continue
                header = _text(ws.cell(max(1, r - 1), c).value) or f"Column {c}"
                if re.search(r"(revenue|eps|earn|ebit|fcf|target|margin|growth)", label, re.I):
                    rows.append({"Timestamp": stamp, "Metric": label, "Period": header, "Value": value, "Source": "Expectations & Consensus sheet"})
    return pd.DataFrame(rows)


def _append_estimates(ticker, snapshot):
    path = _ticker_dir(ticker) / "estimate_history.csv"
    if snapshot.empty:
        return pd.DataFrame()
    if path.exists():
        old = pd.read_csv(path)
        df = pd.concat([old, snapshot], ignore_index=True)
    else:
        df = snapshot.copy()
    df["SnapshotDate"] = df["Timestamp"].astype(str).str[:10]
    df = df.drop_duplicates(["SnapshotDate", "Metric", "Period"], keep="last")
    df.to_csv(path, index=False)
    return df


def _estimate_revisions(history):
    if history.empty:
        return pd.DataFrame()
    rows = []
    h = history.sort_values("Timestamp")
    for (metric, period), part in h.groupby(["Metric", "Period"]):
        if len(part) < 2:
            continue
        prev = part.iloc[-2]; cur = part.iloc[-1]
        old = _num(prev["Value"]); new = _num(cur["Value"])
        if old is None or new is None:
            continue
        change = new - old
        change_pct = change / abs(old) if abs(old) > 1e-12 else np.nan
        rows.append({
            "Metric": metric, "Period": period,
            "PreviousSnapshot": prev["Timestamp"], "CurrentSnapshot": cur["Timestamp"],
            "Previous": old, "Current": new, "Revision": change, "RevisionPct": change_pct,
            "Direction": "UP" if change > 0 else "DOWN" if change < 0 else "UNCHANGED",
            "Source": cur.get("Source"),
        })
    return pd.DataFrame(rows)


def _earnings_history(ticker):
    if yf is None:
        return pd.DataFrame()
    try:
        obj = yf.Ticker(ticker)
        method = getattr(obj, "get_earnings_history", None)
        if not method:
            return pd.DataFrame()
        df = method()
        if df is None or getattr(df, "empty", True):
            return pd.DataFrame()
        out = df.reset_index().copy()
        out.columns = [str(c) for c in out.columns]
        return out.tail(12)
    except Exception:
        return pd.DataFrame()


def ensure_earnings_revisions(wb, ticker, info):
    snapshot = _consensus_snapshot(wb, ticker, info or {})
    history = _append_estimates(ticker, snapshot)
    revisions = _estimate_revisions(history)
    earnings = _earnings_history(ticker)

    ws = _sheet(wb, "Earnings & Revisions")
    _title(ws, f"{ticker} — Earnings & Estimate Revisions", "Tracks how forward expectations change between model builds and records public earnings-surprise history when the data provider exposes it.")
    _section(ws, 5, "Latest Estimate Revisions")
    _header(ws, 6, ["Metric", "Period", "Previous", "Current", "Revision", "Revision %", "Direction", "Source"])
    for rr, (_, row) in enumerate(revisions.tail(30).iterrows(), 7):
        vals = [row["Metric"], row["Period"], row["Previous"], row["Current"], row["Revision"], row["RevisionPct"], row["Direction"], row.get("Source")]
        for c, v in enumerate(vals, 1): ws.cell(rr, c, v)
        ws.cell(rr, 6).number_format = FMT_PCT
    if revisions.empty:
        ws["A7"] = "REVIEW"
        ws["B7"] = "A second point-in-time estimate snapshot is required before revisions can be measured."

    start = max(40, ws.max_row + 3)
    _section(ws, start, "Recent Earnings Surprise History")
    if earnings.empty:
        ws.cell(start + 1, 1, "REVIEW")
        ws.cell(start + 1, 2, "Provider earnings-history endpoint unavailable; no surprise is fabricated.")
    else:
        cols = list(earnings.columns)[:8]
        _header(ws, start + 1, cols)
        for rr, (_, row) in enumerate(earnings.iterrows(), start + 2):
            for c, col in enumerate(cols, 1):
                ws.cell(rr, c, row[col] if not pd.isna(row[col]) else None)

    status = "PASS" if not revisions.empty or not earnings.empty else "REVIEW"
    _quality_row(wb, "Earnings / estimate revision history", status, f"{len(history)} durable estimate snapshots; {len(revisions)} latest revisions; {len(earnings)} earnings-history rows.")
    for col, width in {"A":34, "B":24, "C":18, "D":18, "E":18, "F":15, "G":14, "H":50}.items(): ws.column_dimensions[col].width = width
    ws.freeze_panes = "A7"
    return {"history": history, "revisions": revisions, "earnings": earnings}


def ensure_capital_allocation(wb, ticker, info):
    hist = _historical(wb)
    tax = _num(info.get("effectiveTaxRate"), 0.21) if info else 0.21
    if tax is None or not (0 <= tax <= .60):
        tax = .21
    wacc = _base_wacc(wb)

    repurchases = _statement_series(wb, ["Repurchase of Capital Stock", "Common Stock Repurchased", "Payments for Repurchase of Common Stock"])
    dividends = _statement_series(wb, ["Dividends Paid", "Common Stock Dividends Paid", "Payments of Dividends"])
    acquisitions = _statement_series(wb, ["Payments to Acquire Businesses, Net of Cash Acquired", "Acquisitions", "Business Acquisitions"])
    assets = _statement_series(wb, ["Total Assets"])
    cash = _statement_series(wb, ["Cash and Cash Equivalents", "Cash, Cash Equivalents and Short Term Investments"])
    current_liab = _statement_series(wb, ["Total Current Liabilities", "Current Liabilities"])

    rows = []
    for _, h in hist.iterrows():
        year = int(h["Year"])
        op = _num(h.get("OperatingIncome")); rev = _num(h.get("Revenue")); fcf = _num(h.get("FCF"))
        shares = _num(h.get("DilutedShares")); nopat = op * (1 - tax) if op is not None else None
        invested = None
        if year in assets:
            invested = assets[year] - (cash.get(year) or 0) - (current_liab.get(year) or 0)
            if invested <= 0:
                invested = None
        rows.append({
            "Year": year,
            "Revenue": rev,
            "NOPAT": nopat,
            "FCF": fcf,
            "Capex": _num(h.get("Capex")),
            "SBC": _num(h.get("SBC")),
            "Repurchases": abs(repurchases.get(year)) if _num(repurchases.get(year)) is not None else None,
            "Dividends": abs(dividends.get(year)) if _num(dividends.get(year)) is not None else None,
            "M&A": abs(acquisitions.get(year)) if _num(acquisitions.get(year)) is not None else None,
            "RepurchasesLessSBC": (
                abs(repurchases.get(year)) - (_num(h.get("SBC"), 0.0) or 0.0)
                if _num(repurchases.get(year)) is not None else None
            ),
            "DilutedShares": shares,
            "ApproxInvestedCapital": invested,
            "ROIC": nopat / invested if nopat is not None and invested not in (None, 0) else None,
        })
    df = pd.DataFrame(rows)
    if not df.empty:
        df["NetShareReduction"] = -df["DilutedShares"].pct_change()
    incremental_roic = None
    clean = df.dropna(subset=["NOPAT", "ApproxInvestedCapital"]) if not df.empty else pd.DataFrame()
    if len(clean) >= 2:
        first, last = clean.iloc[0], clean.iloc[-1]
        delta_cap = last["ApproxInvestedCapital"] - first["ApproxInvestedCapital"]
        if abs(delta_cap) > 1e-9:
            incremental_roic = (last["NOPAT"] - first["NOPAT"]) / delta_cap

    ws = _sheet(wb, "Capital Allocation")
    _title(ws, f"{ticker} — Capital Allocation", "Tracks where cash is deployed and whether incremental operating capital appears to earn above the company's cost of capital. Statement-derived approximations are labeled explicitly.")
    _section(ws, 5, "Capital Allocation History")
    _header(ws, 6, ["Year", "FCF", "Capex", "Repurchases", "SBC", "Repurchases − SBC", "Dividends", "M&A", "Net Share Reduction", "ROIC"])
    for rr, (_, row) in enumerate(df.iterrows(), 7):
        vals = [
            int(row["Year"]), row["FCF"], row["Capex"], row["Repurchases"], row["SBC"],
            row["RepurchasesLessSBC"], row["Dividends"], row["M&A"],
            row.get("NetShareReduction"), row.get("ROIC")
        ]
        for c, v in enumerate(vals, 1): ws.cell(rr, c, None if pd.isna(v) else v)
        for c in (2,3,4,5,6,7,8): ws.cell(rr, c).number_format = FMT_BN
        for c in (9,10): ws.cell(rr, c).number_format = FMT_PCT

    start = max(18, ws.max_row + 3)
    _section(ws, start, "Capital-Efficiency Summary")
    summary = [
        ("Calculated WACC", wacc, "Cost of capital from the workbook"),
        ("Incremental ROIC", incremental_roic, "ΔNOPAT / Δapproximate invested capital across available history"),
        ("Incremental ROIC − WACC", incremental_roic - wacc if incremental_roic is not None and wacc is not None else None, "Positive spread suggests incremental capital earned above the model cost of capital"),
        ("Latest net share reduction", _num(df.iloc[-1]["NetShareReduction"]) if not df.empty else None, "Negative values indicate dilution; buyback cash and SBC should be read together"),
    ]
    _header(ws, start + 1, ["Metric", "Value", "Interpretation"])
    for rr, row in enumerate(summary, start + 2):
        ws.cell(rr, 1, row[0]); ws.cell(rr, 2, row[1]); ws.cell(rr, 3, row[2]); ws.cell(rr, 2).number_format = FMT_PCT
        ws.cell(rr, 3).alignment = Alignment(wrap_text=True)

    status = "PASS" if incremental_roic is not None else "REVIEW"
    _quality_row(wb, "Capital-allocation / incremental ROIC coverage", status, "Incremental ROIC is calculated only when comparable NOPAT and invested-capital history is available; otherwise the sheet remains a cash-allocation audit.")
    for col, width in {"A":16, "B":16, "C":16, "D":18, "E":16, "F":20, "G":16, "H":16, "I":20, "J":16}.items(): ws.column_dimensions[col].width = width
    ws.freeze_panes = "A7"
    return {"history": df, "incremental_roic": incremental_roic, "wacc": wacc}


def _price_history(ticker):
    if yf is None:
        return pd.Series(dtype=float)
    try:
        h = yf.Ticker(ticker).history(period="10y", auto_adjust=True)
        if h is None or h.empty or "Close" not in h:
            return pd.Series(dtype=float)
        s = pd.to_numeric(h["Close"], errors="coerce").dropna()
        s.index = pd.to_datetime(s.index).tz_localize(None)
        return s
    except Exception:
        return pd.Series(dtype=float)


def ensure_valuation_history(wb, ticker):
    hist = _historical(wb)
    prices = _price_history(ticker)
    rows = []
    if not hist.empty and not prices.empty:
        for _, h in hist.iterrows():
            y = int(h["Year"])
            year_px = prices[prices.index.year == y]
            if year_px.empty:
                continue
            px = float(year_px.iloc[-1])
            eps = _num(h.get("EPS"))
            shares = _num(h.get("DilutedShares"))
            fcf = _num(h.get("FCF"))
            fcf_ps = fcf / shares if fcf is not None and shares not in (None, 0) else None
            rows.append({
                "Year": y, "YearEndPrice": px,
                "P/E": px / eps if eps not in (None, 0) else None,
                "FCFYield": fcf_ps / px if fcf_ps is not None and px else None,
            })
    df = pd.DataFrame(rows)
    current_price = _current_price(wb)
    latest = hist.iloc[-1] if not hist.empty else None
    current_pe = current_fcf_yield = None
    if latest is not None and current_price:
        eps = _num(latest.get("EPS")); shares = _num(latest.get("DilutedShares")); fcf = _num(latest.get("FCF"))
        current_pe = current_price / eps if eps not in (None, 0) else None
        fcf_ps = fcf / shares if fcf is not None and shares not in (None, 0) else None
        current_fcf_yield = fcf_ps / current_price if fcf_ps is not None else None

    def pctile(series, value):
        vals = pd.to_numeric(series, errors="coerce").dropna().to_numpy()
        if value is None or len(vals) < 3:
            return None
        return float((vals < value).mean())

    pe_pct = pctile(df["P/E"], current_pe) if not df.empty else None
    fy_pct = pctile(df["FCFYield"], current_fcf_yield) if not df.empty else None

    ws = _sheet(wb, "Valuation History")
    _title(ws, f"{ticker} — Historical Valuation Regime", "Year-end price is paired with the corresponding annual EPS and FCF-per-share history. This is contextual valuation history, not a point-in-time consensus backtest.")
    _section(ws, 5, "Historical Multiples")
    _header(ws, 6, ["Year", "Year-End Price", "P/E", "FCF Yield", "P/E vs History", "FCF Yield vs History", "Status"])
    pe_med = df["P/E"].median() if not df.empty else np.nan
    fy_med = df["FCFYield"].median() if not df.empty else np.nan
    for rr, (_, row) in enumerate(df.iterrows(), 7):
        vals = [int(row["Year"]), row["YearEndPrice"], row["P/E"], row["FCFYield"], row["P/E"]/pe_med-1 if pd.notna(pe_med) and pe_med else None, row["FCFYield"]/fy_med-1 if pd.notna(fy_med) and fy_med else None, "PASS"]
        for c, v in enumerate(vals, 1): ws.cell(rr, c, None if pd.isna(v) else v)
        ws.cell(rr, 2).number_format = FMT_PRICE; ws.cell(rr, 3).number_format = FMT_MULT
        for c in (4,5,6): ws.cell(rr, c).number_format = FMT_PCT

    start = max(18, ws.max_row + 3)
    _section(ws, start, "Current Valuation Context")
    _header(ws, start + 1, ["Metric", "Current", "Historical Median", "Current Percentile", "Interpretation"])
    contexts = [
        ("P/E", current_pe, pe_med if pd.notna(pe_med) else None, pe_pct, "Higher percentile = more expensive relative to observed annual history"),
        ("FCF Yield", current_fcf_yield, fy_med if pd.notna(fy_med) else None, fy_pct, "Higher percentile = higher cash yield relative to observed annual history"),
    ]
    for rr, row in enumerate(contexts, start + 2):
        for c, v in enumerate(row, 1): ws.cell(rr, c, v)
        ws.cell(rr, 2).number_format = FMT_MULT if row[0] == "P/E" else FMT_PCT
        ws.cell(rr, 3).number_format = FMT_MULT if row[0] == "P/E" else FMT_PCT
        ws.cell(rr, 4).number_format = FMT_PCT
        ws.cell(rr, 5).alignment = Alignment(wrap_text=True)

    status = "PASS" if len(df) >= 3 else "REVIEW"
    _quality_row(wb, "Historical valuation regime coverage", status, f"{len(df)} annual valuation observations. History uses year-end market prices and annual reported fundamentals; it is not a point-in-time analyst-estimate series.")
    for col, width in {"A":16, "B":20, "C":16, "D":16, "E":20, "F":24, "G":15}.items(): ws.column_dimensions[col].width = width
    return {"history": df, "current_pe": current_pe, "current_fcf_yield": current_fcf_yield, "pe_percentile": pe_pct, "fcf_yield_percentile": fy_pct}


def _segment_candidates(wb):
    if "Segment Analysis" not in wb.sheetnames:
        return []
    ws = wb["Segment Analysis"]
    rows = []
    for r in range(1, min(ws.max_row, 80) + 1):
        name = _text(ws.cell(r, 1).value)
        if not name or len(name) > 80:
            continue
        low = name.lower()
        if any(x in low for x in ["segment analysis", "source", "data quality", "metric", "revenue by", "business line"]):
            continue
        nums = []
        for c in range(2, min(ws.max_column, 10) + 1):
            v = _num(ws.cell(r, c).value)
            if v is not None:
                nums.append(v)
        if nums:
            rows.append((name, nums[-1]))
    # De-duplicate while preserving order.
    seen = set(); out = []
    for name, value in rows:
        key = name.lower()
        if key in seen:
            continue
        seen.add(key); out.append((name, value))
    return out[:12]


def ensure_sotp_framework(wb, ticker):
    segments = _segment_candidates(wb)
    path = _ticker_dir(ticker) / "sotp_assumptions.csv"
    assumptions = pd.DataFrame()
    if path.exists():
        try:
            assumptions = pd.read_csv(path)
        except Exception:
            assumptions = pd.DataFrame()
    if not path.exists() and segments:
        pd.DataFrame({
            "Segment": [x[0] for x in segments],
            "MetricBasis": ["Latest disclosed segment metric"] * len(segments),
            "Multiple": [np.nan] * len(segments),
            "AdjustmentBn": [0.0] * len(segments),
            "Notes": ["Enter an explicit research multiple; blank means REVIEW / no valuation."] * len(segments),
        }).to_csv(path, index=False)

    amap = {}
    if not assumptions.empty and "Segment" in assumptions:
        for _, row in assumptions.iterrows():
            amap[str(row["Segment"]).strip().lower()] = row

    ws = _sheet(wb, "SOTP Framework")
    _title(ws, f"{ticker} — Sum-of-the-Parts Framework", "SOTP is enabled only when segment economics exist and explicit local valuation multiples are supplied. Blank multiples remain REVIEW; no segment value is fabricated.")
    _section(ws, 5, "Segment Valuation")
    _header(ws, 6, ["Segment", "Latest Metric", "Metric Basis", "Multiple", "Adjustment", "Implied Segment Value", "Status", "Notes"])
    total = 0.0; valid = 0
    for rr, (segment, metric) in enumerate(segments, 7):
        a = amap.get(segment.lower())
        multiple = _num(a.get("Multiple")) if a is not None else None
        adjustment = _num(a.get("AdjustmentBn"), 0.0) if a is not None else 0.0
        basis = _text(a.get("MetricBasis")) if a is not None else "Latest disclosed segment metric"
        notes = _text(a.get("Notes")) if a is not None else f"Populate {path.name} locally to activate SOTP."
        value = metric * multiple + adjustment if multiple is not None else None
        status = "PASS" if value is not None else "REVIEW"
        if value is not None:
            total += value; valid += 1
        vals = [segment, metric, basis, multiple, adjustment, value, status, notes]
        for c, v in enumerate(vals, 1): ws.cell(rr, c, v)
        ws.cell(rr, 2).number_format = FMT_BN; ws.cell(rr, 4).number_format = FMT_MULT
        ws.cell(rr, 5).number_format = FMT_BN; ws.cell(rr, 6).number_format = FMT_BN
        ws.cell(rr, 7).fill = _status_fill(status)

    start = max(24, ws.max_row + 3)
    _section(ws, start, "Equity-Value Bridge")
    shares = _shares(wb); net_debt = _net_debt(wb)
    equity = total - net_debt if valid and net_debt is not None else total if valid else None
    per_share = equity / shares if equity is not None and shares not in (None, 0) else None
    summary = [
        ("Valued segment total", total if valid else None, "Sum of segments with explicit multiples"),
        ("Net debt / (net cash)", net_debt, "Company Data; negative net debt increases equity value"),
        ("SOTP equity value", equity, "Segment total less net debt"),
        ("SOTP value / share", per_share, "Requires diluted shares"),
        ("Assumption file", _display_path(path), "Local/private input; ignored by Git"),
    ]
    _header(ws, start + 1, ["Metric", "Value", "Interpretation"])
    for rr, row in enumerate(summary, start + 2):
        for c, v in enumerate(row, 1): ws.cell(rr, c, v)
        ws.cell(rr, 2).number_format = FMT_PRICE if row[0] == "SOTP value / share" else FMT_BN
        ws.cell(rr, 3).alignment = Alignment(wrap_text=True)

    status = "PASS" if valid >= 2 else "REVIEW"
    _quality_row(wb, "SOTP valuation reliability gate", status, f"{valid} segment(s) have explicit local multiples out of {len(segments)} detected segment rows. SOTP is not used as evidence until sufficient segment economics and assumptions exist.")
    for col, width in {"A":34, "B":18, "C":32, "D":14, "E":16, "F":22, "G":14, "H":55}.items(): ws.column_dimensions[col].width = width
    return {"segments": segments, "valid": valid, "value_per_share": per_share, "path": path}


def ensure_thesis_timeline(wb, ticker):
    path = _ticker_dir(ticker) / "thesis_timeline.csv"
    columns = ["Date", "EventType", "Event", "Expected", "Actual", "ThesisImpact", "Evidence", "Source"]
    if path.exists():
        try:
            df = pd.read_csv(path)
        except Exception:
            df = pd.DataFrame(columns=columns)
    else:
        df = pd.DataFrame(columns=columns)

    decision = _decision_values(wb)
    today = datetime.now().astimezone().date().isoformat()
    event = f"Model build: price={decision.get('CurrentPrice')}; base value={decision.get('BaseValue')}; view={decision.get('ModelView') or 'N/A'}"
    already = (
        not df.empty
        and ((df.get("Date", pd.Series(dtype=str)).astype(str) == today)
             & (df.get("EventType", pd.Series(dtype=str)).astype(str) == "MODEL SNAPSHOT")).any()
    )
    if not already:
        row = pd.DataFrame([{
            "Date": today, "EventType": "MODEL SNAPSHOT", "Event": event,
            "Expected": "", "Actual": "", "ThesisImpact": "NEUTRAL",
            "Evidence": "Deterministic workbook snapshot", "Source": "Local model build",
        }])
        df = pd.concat([df, row], ignore_index=True)
        df.to_csv(path, index=False)

    ws = _sheet(wb, "Thesis Timeline")
    _title(ws, f"{ticker} — Thesis & Catalyst Timeline", "Chronological evidence log. Add earnings, catalysts, regulatory events and falsification tests to the local CSV; model snapshots are appended automatically once per day.")
    _section(ws, 5, "Thesis Events")
    _header(ws, 6, columns)
    view = df.tail(30) if not df.empty else df
    for rr, (_, row) in enumerate(view.iterrows(), 7):
        for c, col in enumerate(columns, 1):
            ws.cell(rr, c, row.get(col))
            ws.cell(rr, c).alignment = Alignment(wrap_text=True, vertical="top")
    ws["A39"] = "Local input"
    ws["B39"] = _display_path(path)
    ws["B39"].font = Font(color="008000")
    _quality_row(wb, "Thesis / catalyst chronology", "PASS" if len(df) >= 1 else "REVIEW", f"{len(df)} durable thesis event(s) stored locally. Manual events should include expected vs actual evidence and thesis impact.")
    for col, width in {"A":15, "B":18, "C":48, "D":34, "E":34, "F":18, "G":50, "H":45}.items(): ws.column_dimensions[col].width = width
    ws.freeze_panes = "A7"
    return df


def apply_research_accountability(wb, ticker, info=None):
    """Apply all offline/private research-accountability layers to an in-memory workbook."""
    info = info or {}
    results = {}
    for name, fn, args in [
        ("forecast", ensure_forecast_accountability, (wb, ticker)),
        ("earnings", ensure_earnings_revisions, (wb, ticker, info)),
        ("capital_allocation", ensure_capital_allocation, (wb, ticker, info)),
        ("valuation_history", ensure_valuation_history, (wb, ticker)),
        ("sotp", ensure_sotp_framework, (wb, ticker)),
        ("thesis_timeline", ensure_thesis_timeline, (wb, ticker)),
    ]:
        try:
            results[name] = fn(*args)
        except Exception as exc:
            results[name] = {"status": "ERROR", "error": str(exc)}
            _quality_row(wb, f"{name.replace('_',' ').title()} extension", "REVIEW", f"Extension could not complete: {exc}")
    return results
