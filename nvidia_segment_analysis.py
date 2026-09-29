from __future__ import annotations

"""Verified NVIDIA segment adapter using issuer-reported SEC 10-K values."""

from openpyxl.styles import Alignment, Font, PatternFill

NAVY="17365D"; BLUE="2F75B5"; WHITE="FFFFFF"; GREY="666666"; LINK_GREEN="008000"
FMT_BN='#,##0.0;[Red](#,##0.0);-'; FMT_PCT='0.0%;[Red](0.0%);-'
SOURCE="https://www.sec.gov/Archives/edgar/data/1045810/000104581026000021/nvda-20260125.htm"
SEGMENTS={
    "Compute & Networking":{"revenue":[47.405,116.193,193.479],"profit":[32.016,82.875,130.141]},
    "Graphics":{"revenue":[13.517,14.304,22.459],"profit":[5.846,5.085,9.156]},
}
BUSINESS={
    "Data Center":[47.525,115.186,193.737],
    "Gaming":[10.447,11.350,16.042],
    "Professional Visualization":[1.553,1.878,3.191],
    "Automotive":[1.091,1.694,2.349],
    "OEM and Other":[0.306,0.389,0.619],
}

def _fill(c): return PatternFill("solid",fgColor=c)

def ensure_nvidia_segment_analysis(wb,ticker="NVDA"):
    if str(ticker or "").upper()!="NVDA": return False
    if "Segment Analysis" in wb.sheetnames: wb.remove(wb["Segment Analysis"])
    ws=wb.create_sheet("Segment Analysis"); ws.sheet_view.showGridLines=False
    for c in range(1,17): ws.cell(1,c).fill=_fill(NAVY)
    ws["A1"]="NVDA — Business & Segment Analysis"; ws["A1"].font=Font(bold=True,color=WHITE,size=18)
    ws["A3"]="Verified NVIDIA adapter. Fiscal 2024–2026 reportable-segment revenue and operating income, plus end-market revenue, are taken directly from NVIDIA's fiscal 2026 Form 10-K."
    ws.merge_cells("A3:P3"); ws["A3"].font=Font(italic=True,color=GREY); ws["A3"].alignment=Alignment(wrap_text=True)
    for c in range(1,17): ws.cell(5,c).fill=_fill(NAVY); ws.cell(5,c).font=Font(bold=True,color=WHITE)
    ws["A5"]="Reported Operating / Reportable Segments"
    heads=["Segment","2024 Revenue","2025 Revenue","2026 Revenue","2026 Growth","2024–2026 CAGR","2024 Segment Profit","2025 Segment Profit","2026 Segment Profit","2024 Margin","2025 Margin","2026 Margin","Margin Δ","2026 Revenue Mix","Data Status","Source / Notes"]
    for c,h in enumerate(heads,1): ws.cell(6,c,h); ws.cell(6,c).fill=_fill(BLUE); ws.cell(6,c).font=Font(bold=True,color=WHITE)
    for r,(name,d) in enumerate(SEGMENTS.items(),7):
        vals=[name,*d["revenue"],None,None,*d["profit"],None,None,None,None,None,"Verified revenue + operating income",SOURCE]
        for c,v in enumerate(vals,1): ws.cell(r,c,v)
        ws.cell(r,5,f'=D{r}/C{r}-1'); ws.cell(r,6,f'=(D{r}/B{r})^(1/2)-1')
        ws.cell(r,10,f'=G{r}/B{r}'); ws.cell(r,11,f'=H{r}/C{r}'); ws.cell(r,12,f'=I{r}/D{r}')
        ws.cell(r,13,f'=L{r}-K{r}'); ws.cell(r,14,f'=D{r}/SUM($D$7:$D$8)')
        for c in (2,3,4,7,8,9): ws.cell(r,c).number_format=FMT_BN
        for c in (5,6,10,11,12,13,14): ws.cell(r,c).number_format=FMT_PCT
        ws.cell(r,16).hyperlink=SOURCE; ws.cell(r,16).font=Font(color=LINK_GREEN,underline="single")
    for c in range(1,10): ws.cell(11,c).fill=_fill(NAVY); ws.cell(11,c).font=Font(bold=True,color=WHITE)
    ws["A11"]="Revenue by End Market"
    bh=["End Market","2024","2025","2026","2026 Growth","2024–2026 CAGR","2026 Mix","Source / Notes"]
    for c,h in enumerate(bh,1): ws.cell(12,c,h); ws.cell(12,c).fill=_fill(BLUE); ws.cell(12,c).font=Font(bold=True,color=WHITE)
    for r,(name,rev) in enumerate(BUSINESS.items(),13):
        for c,v in enumerate([name,*rev,None,None,None,SOURCE],1): ws.cell(r,c,v)
        ws.cell(r,5,f'=D{r}/C{r}-1'); ws.cell(r,6,f'=(D{r}/B{r})^(1/2)-1'); ws.cell(r,7,f'=D{r}/SUM($D$13:$D$17)')
        for c in (2,3,4): ws.cell(r,c).number_format=FMT_BN
        for c in (5,6,7): ws.cell(r,c).number_format=FMT_PCT
        ws.cell(r,8).hyperlink=SOURCE; ws.cell(r,8).font=Font(color=LINK_GREEN,underline="single")
    for c in range(1,10): ws.cell(20,c).fill=_fill(NAVY); ws.cell(20,c).font=Font(bold=True,color=WHITE)
    ws["A20"]="Source & Data Quality"
    ws["A21"]="SEC 10-K Source"; ws["B21"]=SOURCE; ws["B21"].hyperlink=SOURCE; ws["B21"].font=Font(color=LINK_GREEN,underline="single")
    ws["A22"]="Extraction Status"; ws["B22"]="VERIFIED ADAPTER — NVIDIA fiscal 2026 Form 10-K"
    ws["A23"]="Important"; ws["B23"]="Segment operating income differs from consolidated operating income because certain enterprise expenses are not allocated to reportable segments."
    ws["B23"].alignment=Alignment(wrap_text=True)
    widths={"A":34,"B":16,"C":16,"D":16,"E":14,"F":16,"G":20,"H":20,"I":20,"J":14,"K":14,"L":14,"M":14,"N":16,"O":28,"P":60}
    for col,w in widths.items(): ws.column_dimensions[col].width=w
    ws.freeze_panes="A7"
    return True
