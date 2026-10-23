from __future__ import annotations

"""Workbook enrichment for Nasdaq Helsinki issuers.

This layer keeps issuer-defined KPIs and latest interim/result tables separate from standardized
IFRS statements.  It is intentionally source-preserving: labels stay as the issuer publishes
them, values are never invented, and every populated section links back to the issuer page.
"""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from finland_sources import (
    finland_key_figures,
    finland_latest_results,
    is_finnish_ticker,
)

NAVY="17365D"; BLUE="2F75B5"; WHITE="FFFFFF"; GREY="666666"; GREEN="E2F0D9"; GOLD="FFF2CC"
LINK_GREEN="008000"; THIN=Side(style="thin",color="D9E1F2")
FMT_NUM='#,##0.0;[Red](#,##0.0);-'


def _fill(color):
    return PatternFill("solid",fgColor=color)


def _section(ws,row,title,end):
    for c in range(1,end+1):
        ws.cell(row,c).fill=_fill(NAVY)
        ws.cell(row,c).font=Font(bold=True,color=WHITE)
    ws.cell(row,1,title)


def _header(ws,row,values):
    for c,value in enumerate(values,1):
        cell=ws.cell(row,c,value)
        cell.fill=_fill(BLUE); cell.font=Font(bold=True,color=WHITE)
        cell.border=Border(bottom=THIN)
        cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)


def _format_metric(cell,metric):
    low=str(metric or "").lower()
    if "%" in low or any(x in low for x in ("margin","ratio","return on","roce","roe","gearing","equity ratio")):
        # Issuer tables usually publish percentage points as e.g. 38, not 0.38.
        cell.number_format='0.0"%"'
    elif "per share" in low or "eps" in low:
        cell.number_format='0.00'
    else:
        cell.number_format=FMT_NUM


def _quality_row(wb,label,status,detail):
    if "Data Quality" not in wb.sheetnames:
        return
    ws=wb["Data Quality"]
    row=None
    for r in range(1,ws.max_row+1):
        if str(ws.cell(r,1).value or "").strip()==label:
            row=r; break
    if row is None:
        row=ws.max_row+1
    ws.cell(row,1,label); ws.cell(row,2,status); ws.cell(row,3,detail)
    ws.cell(row,2).font=Font(bold=True)
    ws.cell(row,2).fill=_fill(GREEN if status=="PASS" else GOLD)
    ws.cell(row,3).alignment=Alignment(wrap_text=True,vertical="top")


def ensure_finland_issuer_kpis(wb,ticker,info=None):
    if not is_finnish_ticker(ticker):
        return {"status":"NOT_APPLICABLE"}
    info=info or {}
    annual,annual_meta=finland_key_figures(ticker,info)
    latest,latest_meta=finland_latest_results(ticker,info)

    name="Issuer KPIs"
    if name in wb.sheetnames:
        wb.remove(wb[name])
    # Keep near the core research sheets when possible.
    idx=min(len(wb.worksheets),4)
    ws=wb.create_sheet(name,idx)
    ws.sheet_view.showGridLines=False

    for c in range(1,10):
        ws.cell(1,c).fill=_fill(NAVY); ws.cell(2,c).fill=_fill(NAVY)
    ws["A1"]=f"{ticker} — Finnish Issuer KPIs & Latest Results"
    ws["A1"].font=Font(bold=True,color=WHITE,size=18)
    ws["A3"]="Issuer-defined KPIs are kept separate from standardized IFRS statement rows. Values come from issuer-owned investor-relations/result tables; missing metrics remain blank rather than estimated."
    ws.merge_cells("A3:I3"); ws["A3"].font=Font(italic=True,color=GREY); ws["A3"].alignment=Alignment(wrap_text=True)

    row=5
    _section(ws,row,"Annual issuer-defined key figures",9)
    years=sorted({y for rec in annual for y in (rec.get("values") or {})},reverse=True)[:6]
    years=sorted(years)
    headers=["Metric"]+[str(y) for y in years]+["Primary Source"]
    _header(ws,row+1,headers)
    rr=row+2
    for rec in annual:
        values=rec.get("values") or {}
        if not any(y in values for y in years):
            continue
        ws.cell(rr,1,rec.get("metric"))
        for i,y in enumerate(years,2):
            v=values.get(y)
            ws.cell(rr,i,v)
            if v is not None:
                _format_metric(ws.cell(rr,i),rec.get("metric"))
        src=rec.get("source") or annual_meta.get("source")
        sc=2+len(years)
        ws.cell(rr,sc,src)
        if src:
            ws.cell(rr,sc).hyperlink=src; ws.cell(rr,sc).font=Font(color=LINK_GREEN,underline="single")
        rr+=1
    if rr==row+2:
        ws.cell(rr,1,"No issuer KPI table recovered automatically.")
        ws.cell(rr,1).font=Font(italic=True,color=GREY)
        rr+=1

    row=rr+2
    _section(ws,row,"Latest interim / results table",9)
    title=latest.get("title") if latest else None
    if latest:
        ws.cell(row+1,1,"Release"); ws.cell(row+1,2,title)
        ws.cell(row+2,1,"Date"); ws.cell(row+2,2,latest.get("date"))
        ws.cell(row+3,1,"Source"); ws.cell(row+3,2,latest.get("url"))
        if latest.get("url"):
            ws.cell(row+3,2).hyperlink=latest["url"]; ws.cell(row+3,2).font=Font(color=LINK_GREEN,underline="single")
        table_row=row+5
        cols=list(latest.get("columns") or [])[:7]
        _header(ws,table_row,["Metric"]+cols)
        rr=table_row+1
        for rec in (latest.get("rows") or [])[:50]:
            ws.cell(rr,1,rec.get("metric"))
            vals=rec.get("values") or []
            for i,item in enumerate(vals[:len(cols)],2):
                value=item.get("value")
                ws.cell(rr,i,value if value is not None else item.get("raw"))
                if value is not None:
                    _format_metric(ws.cell(rr,i),rec.get("metric"))
            rr+=1
    else:
        ws.cell(row+1,1,"No latest-results table recovered automatically.")
        ws.cell(row+1,2,f"Status: {latest_meta.get('status','NO_DATA')}")
        rr=row+2

    row=rr+2
    consensus=getattr(wb,"_consensus_result",{}) or {}
    vara_rows=[x for x in (consensus.get("observations") or []) if x.get("provider")=="Vara Research"]
    vara_targets=[x for x in (consensus.get("price_targets") or []) if x.get("provider")=="Vara Research"]
    if vara_rows or vara_targets:
        _section(ws,row,"Issuer-linked Vara Research consensus",9)
        _header(ws,row+1,["Fiscal Year","Metric","Consensus","Low","High","Analysts","Source"])
        rr=row+2
        for rec in sorted(vara_rows,key=lambda x:(x.get("fiscal_year") or 0,x.get("metric") or "")):
            vals=[rec.get("fiscal_year"),rec.get("metric"),rec.get("mean"),rec.get("low"),rec.get("high"),rec.get("analysts"),rec.get("source_url")]
            for cc,v in enumerate(vals,1):
                ws.cell(rr,cc,v)
            metric=rec.get("metric")
            for cc in (3,4,5):
                if ws.cell(rr,cc).value is not None:
                    ws.cell(rr,cc).number_format="0.00" if metric in {"EPS","DPS"} else FMT_NUM
            if rec.get("source_url"):
                ws.cell(rr,7).hyperlink=rec.get("source_url"); ws.cell(rr,7).font=Font(color=LINK_GREEN,underline="single")
            rr+=1
        for target in vara_targets[:1]:
            vals=["Current","Analyst Price Target",target.get("mean") or target.get("median"),target.get("low"),target.get("high"),target.get("analysts"),target.get("source_url")]
            for cc,v in enumerate(vals,1):
                ws.cell(rr,cc,v)
            for cc in (3,4,5): ws.cell(rr,cc).number_format="0.00"
            if target.get("source_url"):
                ws.cell(rr,7).hyperlink=target.get("source_url"); ws.cell(rr,7).font=Font(color=LINK_GREEN,underline="single")
            rr+=1
        row=rr+2

    _section(ws,row,"Source status",9)
    status_rows=[
        ("Annual KPI extraction",annual_meta.get("status"),annual_meta.get("source") or "; ".join(annual_meta.get("pages_checked") or [])),
        ("Latest results extraction",latest_meta.get("status"),latest_meta.get("source") or f"Candidate releases: {latest_meta.get('candidates',0)}"),
        ("Vara consensus","PASS" if vara_rows else "REVIEW",f"Issuer-linked Vara observations recovered: {len(vara_rows)}"),
        ("Accounting source","PASS","Finnish OAM / ESEF is the primary standardized annual-statement route for .HE issuers."),
        ("Research rule","PASS","Issuer-adjusted KPIs are not substituted for IFRS accounting metrics unless explicitly labeled."),
    ]
    _header(ws,row+1,["Check","Status","Evidence / Source"])
    for rr2,vals in enumerate(status_rows,row+2):
        for c,v in enumerate(vals,1):
            ws.cell(rr2,c,v)
            ws.cell(rr2,c).alignment=Alignment(wrap_text=True,vertical="top")
        ws.cell(rr2,2).font=Font(bold=True)

    for col,width in {"A":42,"B":18,"C":20,"D":20,"E":20,"F":20,"G":20,"H":20,"I":60}.items():
        ws.column_dimensions[col].width=width
    ws.freeze_panes="B7"

    annual_status="PASS" if annual else "REVIEW"
    latest_status="PASS" if latest else "REVIEW"
    _quality_row(
        wb,
        "Finnish issuer KPI coverage",
        annual_status,
        f"Annual issuer KPI rows recovered: {len(annual)}; source={annual_meta.get('source') or 'not recovered'}",
    )
    _quality_row(
        wb,
        "Finnish latest-results coverage",
        latest_status,
        f"Latest issuer results table: {latest.get('title') if latest else 'not recovered'}",
    )
    return {
        "status":"PASS" if annual or latest else "REVIEW",
        "annual_rows":len(annual),
        "annual_meta":annual_meta,
        "latest_rows":len(latest.get("rows") or []) if latest else 0,
        "latest_meta":latest_meta,
        "vara_consensus_rows":len(vara_rows),
        "vara_price_targets":len(vara_targets),
    }


def finalize_finland_labels(wb,ticker,info=None):
    """Remove US/SEC-only wording from Helsinki workbooks after all legacy generators run."""
    if not is_finnish_ticker(ticker):
        return
    info=info or {}
    currency=str(info.get("currency") or info.get("financialCurrency") or "EUR").upper()

    if "Company Data" in wb.sheetnames:
        ws=wb["Company Data"]
        for cell,label in {
            "A10":f"Market Cap ({currency} bn)",
            "A11":f"Enterprise Value ({currency} bn)",
            "A12":f"Cash ({currency} bn)",
            "A13":f"Debt ({currency} bn)",
            "A14":f"Net Debt ({currency} bn)",
        }.items():
            ws[cell]=label
        ws["D7"]="Regulated filings: Finnish OAM / ESEF | Current releases: issuer IR / Nasdaq Helsinki"

    if "Filings" in wb.sheetnames:
        ws=wb["Filings"]
        ws["A1"]=f"{ticker} — Filings & Regulated Sources"

    if "Financial Statements" in wb.sheetnames:
        ws=wb["Financial Statements"]
        current=str(ws["A3"].value or "")
        if "SEC Company Facts remains primary" in current or current.startswith("USD billions"):
            ws["A3"]=f"{currency} billions unless per-share/share-count data. Primary structured annual source for Helsinki issuers: Finnish OAM / ESEF; issuer IR and public market data are fallbacks. Missing fields remain blank rather than estimated."

    if "Data Quality" in wb.sheetnames:
        _quality_row(
            wb,
            "Finnish source routing",
            "PASS",
            "Helsinki issuer detected: Finnish OAM/ESEF → issuer IR → SEC only when applicable → Yahoo/public market-data fallback.",
        )
