from __future__ import annotations

"""Offline/private portfolio decision-loop analytics.

This module links the equity-research workbooks to the institutional portfolio layer while
keeping accounting deterministic and evidence-gated.

Key design rules:
- transaction history is never inferred from current holdings;
- realized TWR/MWR are calculated only when external cash-flow history is present;
- point-in-time weights are reconstructed from the ledger and raw market closes;
- dividends/splits use provider corporate actions unless the ledger explicitly records the
  same event;
- sector allocation/selection attribution requires point-in-time benchmark sector history;
- equity-model expected returns are confidence-shrunk before sizing;
- transaction-cost estimates are explicit diagnostics, not execution instructions.
"""

from datetime import datetime
from pathlib import Path
import math
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook

try:
    import yfinance as yf
except Exception:
    yf = None


LEDGER_COLUMNS = [
    "Date","Ticker","Action","Shares","Price","Fees","CashFlow","Currency",
    "DecisionID","ReferencePrice","SplitRatio","Notes",
]
VALID_ACTIONS = {"BUY","SELL","DEPOSIT","WITHDRAW","DIVIDEND","FEE","SPLIT"}


def _num(value, default=None):
    try:
        if isinstance(value, bool) or value in (None, ""):
            return default
        x=float(value)
        return x if math.isfinite(x) else default
    except Exception:
        return default


def load_transaction_ledger(path: str | Path) -> pd.DataFrame:
    path=Path(path)
    if not path.exists():
        return pd.DataFrame(columns=LEDGER_COLUMNS)
    df=pd.read_csv(path, comment="#")
    if df.empty:
        return pd.DataFrame(columns=LEDGER_COLUMNS)
    for col in LEDGER_COLUMNS:
        if col not in df.columns:
            df[col]=np.nan
    df=df[LEDGER_COLUMNS].copy()
    df["Date"]=pd.to_datetime(df["Date"], errors="coerce")
    df["Ticker"]=df["Ticker"].fillna("").astype(str).str.upper().str.strip()
    df["Action"]=df["Action"].fillna("").astype(str).str.upper().str.strip()
    for col in ["Shares","Price","Fees","CashFlow","ReferencePrice","SplitRatio"]:
        df[col]=pd.to_numeric(df[col], errors="coerce")
    df["Fees"]=df["Fees"].fillna(0.0)
    bad=df[~df["Action"].isin(VALID_ACTIONS) | df["Date"].isna()]
    if not bad.empty:
        raise ValueError("Transaction ledger contains invalid dates/actions: "+", ".join(sorted(set(bad["Action"].astype(str)))))
    security_actions=df["Action"].isin({"BUY","SELL","DIVIDEND","SPLIT"})
    if (security_actions & df["Ticker"].eq("")).any():
        raise ValueError("BUY/SELL/DIVIDEND/SPLIT rows require a Ticker.")
    return df.sort_values(["Date","Action","Ticker"]).reset_index(drop=True)


def ledger_tickers(ledger: pd.DataFrame) -> list[str]:
    if ledger is None or ledger.empty:
        return []
    return list(dict.fromkeys(x for x in ledger["Ticker"].astype(str).str.upper() if x and x!="NAN"))


def download_realized_market_data(tickers: list[str], period="max"):
    """Download raw closes + corporate actions for transaction accounting.

    Existing portfolio analytics continue to use adjusted prices separately.
    """
    prices={}
    actions={}
    warnings=[]
    if yf is None:
        return pd.DataFrame(), {}, ["yfinance unavailable"]
    for ticker in dict.fromkeys(str(x).upper() for x in tickers if str(x).strip()):
        try:
            h=yf.Ticker(ticker).history(period=period, auto_adjust=False, actions=True)
            if h is None or h.empty or "Close" not in h:
                warnings.append(f"{ticker}: no raw price history")
                continue
            idx=pd.to_datetime(h.index).tz_localize(None)
            prices[ticker]=pd.Series(pd.to_numeric(h["Close"],errors="coerce").values,index=idx,name=ticker)
            a=pd.DataFrame(index=idx)
            a["Dividends"]=pd.to_numeric(h.get("Dividends",0),errors="coerce").fillna(0.0)
            a["Stock Splits"]=pd.to_numeric(h.get("Stock Splits",0),errors="coerce").fillna(0.0)
            actions[ticker]=a
        except Exception as exc:
            warnings.append(f"{ticker}: {exc}")
    panel=pd.concat(prices.values(),axis=1).sort_index() if prices else pd.DataFrame()
    return panel,actions,warnings


def _next_trading_date(index: pd.DatetimeIndex, dt: pd.Timestamp):
    pos=index.searchsorted(pd.Timestamp(dt), side="left")
    return index[pos] if pos < len(index) else None


def _xirr(cashflows: list[tuple[pd.Timestamp,float]]):
    flows=[(pd.Timestamp(d),float(v)) for d,v in cashflows if _num(v) is not None and abs(float(v))>1e-12]
    if len(flows)<2:
        return None
    start=min(d for d,_ in flows)
    def npv(rate):
        if rate<=-0.999999:
            return float("inf")
        return sum(v/((1+rate)**(((d-start).days)/365.25)) for d,v in flows)
    lo,hi=-.999,10.0
    flo,fhi=npv(lo),npv(hi)
    tries=0
    while flo*fhi>0 and tries<8:
        hi*=2; fhi=npv(hi); tries+=1
    if flo*fhi>0:
        return None
    for _ in range(120):
        mid=(lo+hi)/2; fm=npv(mid)
        if abs(fm)<1e-10:
            return mid
        if flo*fm<=0:
            hi=mid; fhi=fm
        else:
            lo=mid; flo=fm
    return (lo+hi)/2


def build_realized_portfolio(
    ledger: pd.DataFrame,
    raw_prices: pd.DataFrame,
    adjusted_benchmark: pd.Series | None = None,
    corporate_actions: dict[str,pd.DataFrame] | None = None,
) -> dict:
    """Reconstruct daily positions/cash/NAV and realized performance."""
    empty={
        "summary":{"status":"NO_LEDGER"},
        "timeseries":pd.DataFrame(),"weights":pd.DataFrame(),"attribution":pd.DataFrame(),
        "transactions":ledger.copy() if ledger is not None else pd.DataFrame(),"warnings":[],
    }
    if ledger is None or ledger.empty:
        return empty
    if raw_prices is None or raw_prices.empty:
        empty["summary"]={"status":"NO_RAW_PRICES"}
        return empty

    prices=raw_prices.copy()
    prices.index=pd.to_datetime(prices.index).tz_localize(None)
    prices=prices.sort_index()
    tickers=ledger_tickers(ledger)
    missing=[t for t in tickers if t not in prices.columns]
    if missing:
        empty["summary"]={"status":"MISSING_PRICES","missing_tickers":missing}
        return empty

    start=pd.Timestamp(ledger["Date"].min()).normalize()
    idx=prices.index[prices.index>=start]
    if len(idx)<2:
        empty["summary"]={"status":"INSUFFICIENT_HISTORY"}
        return empty
    prices=prices.reindex(idx).ffill()
    corporate_actions=corporate_actions or {}

    # Map ledger events to the next available market date; this is explicit in output.
    events=ledger.copy()
    effective=[]
    for d in events["Date"]:
        effective.append(_next_trading_date(idx,pd.Timestamp(d).normalize()))
    events["EffectiveDate"]=effective
    dropped=events[events["EffectiveDate"].isna()]
    events=events.dropna(subset=["EffectiveDate"]).copy()
    warnings=[]
    if not dropped.empty:
        warnings.append(f"{len(dropped)} ledger event(s) occur after available market history.")
    shifted=(events["EffectiveDate"].dt.normalize()!=events["Date"].dt.normalize()).sum()
    if shifted:
        warnings.append(f"{int(shifted)} non-trading-date event(s) mapped to the next trading date.")

    positions={t:0.0 for t in tickers}
    cash=0.0
    rows=[]
    weight_rows=[]
    attr_rows=[]
    prev_nav=None
    prev_values={t:0.0 for t in tickers}
    prev_prices=None
    have_external=False
    external_cashflows=[]
    transaction_notional=0.0
    implementation_shortfall=0.0

    event_groups={d:g for d,g in events.groupby("EffectiveDate")}
    manual_div_keys=set()
    manual_split_keys=set()
    for _,e in events.iterrows():
        key=(pd.Timestamp(e["EffectiveDate"]),str(e["Ticker"]))
        if e["Action"]=="DIVIDEND": manual_div_keys.add(key)
        if e["Action"]=="SPLIT": manual_split_keys.add(key)

    for dt in idx:
        px=prices.loc[dt]
        external=0.0
        dividends_today={t:0.0 for t in tickers}

        # Provider corporate actions happen before close valuation.
        for t in tickers:
            a=corporate_actions.get(t)
            if a is None or a.empty or dt not in a.index:
                continue
            split=_num(a.loc[dt,"Stock Splits"] if "Stock Splits" in a else None,0.0) or 0.0
            div=_num(a.loc[dt,"Dividends"] if "Dividends" in a else None,0.0) or 0.0
            if split and split!=1 and (dt,t) not in manual_split_keys:
                positions[t]*=split
            if div and positions[t] and (dt,t) not in manual_div_keys:
                amount=positions[t]*div
                cash+=amount; dividends_today[t]+=amount

        if dt in event_groups:
            for _,e in event_groups[dt].iterrows():
                action=e["Action"]; ticker=e["Ticker"]
                shares=_num(e["Shares"],0.0) or 0.0
                fee=_num(e["Fees"],0.0) or 0.0
                trade_px=_num(e["Price"])
                if trade_px is None and ticker in tickers:
                    trade_px=_num(px.get(ticker))
                    warnings.append(f"{dt.date()} {ticker} {action}: missing execution price; close used.")
                if action=="DEPOSIT":
                    flow=abs(_num(e["CashFlow"],0.0) or 0.0)
                    cash+=flow; external+=flow; have_external=True
                    external_cashflows.append((dt,-flow))
                elif action=="WITHDRAW":
                    flow=abs(_num(e["CashFlow"],0.0) or 0.0)
                    cash-=flow; external-=flow; have_external=True
                    external_cashflows.append((dt,flow))
                elif action=="BUY":
                    positions[ticker]+=shares
                    notional=shares*trade_px
                    cash-=notional+fee; transaction_notional+=abs(notional)
                    ref=_num(e["ReferencePrice"])
                    if ref is not None:
                        implementation_shortfall+=(trade_px-ref)*shares+fee
                elif action=="SELL":
                    positions[ticker]-=shares
                    notional=shares*trade_px
                    cash+=notional-fee; transaction_notional+=abs(notional)
                    ref=_num(e["ReferencePrice"])
                    if ref is not None:
                        implementation_shortfall+=(ref-trade_px)*shares+fee
                elif action=="DIVIDEND":
                    flow=_num(e["CashFlow"])
                    if flow is None:
                        flow=shares if shares else 0.0
                    cash+=flow; dividends_today[ticker]+=flow
                elif action=="FEE":
                    flow=abs(_num(e["CashFlow"],fee) or fee)
                    cash-=flow
                elif action=="SPLIT":
                    ratio=_num(e["SplitRatio"],_num(e["Shares"]))
                    if ratio and ratio>0:
                        positions[ticker]*=ratio

        security_values={t:positions[t]*(_num(px.get(t),0.0) or 0.0) for t in tickers}
        nav=cash+sum(security_values.values())
        daily_return=None
        if prev_nav is not None and abs(prev_nav)>1e-12 and have_external:
            daily_return=(nav-external)/prev_nav-1
        row={"Date":dt,"NAV":nav,"Cash":cash,"ExternalFlow":external,"PortfolioReturn":daily_return}
        rows.append(row)

        denom=nav if abs(nav)>1e-12 else np.nan
        for t in tickers:
            weight_rows.append({"Date":dt,"Ticker":t,"Shares":positions[t],"MarketValue":security_values[t],"Weight":security_values[t]/denom if pd.notna(denom) else np.nan})
        weight_rows.append({"Date":dt,"Ticker":"CASH","Shares":np.nan,"MarketValue":cash,"Weight":cash/denom if pd.notna(denom) else np.nan})

        if prev_nav is not None and abs(prev_nav)>1e-12 and prev_prices is not None:
            total_contrib=0.0
            for t in tickers:
                p0=_num(prev_prices.get(t)); p1=_num(px.get(t))
                price_ret=(p1/p0-1) if p0 not in (None,0) and p1 is not None else 0.0
                price_contrib=(prev_values.get(t,0.0)/prev_nav)*price_ret
                div_contrib=dividends_today[t]/prev_nav
                total=price_contrib+div_contrib
                total_contrib+=total
                attr_rows.append({"Date":dt,"Ticker":t,"PriceContribution":price_contrib,"DividendContribution":div_contrib,"TotalContribution":total})
            if daily_return is not None:
                attr_rows.append({"Date":dt,"Ticker":"Trading / cash residual","PriceContribution":0.0,"DividendContribution":0.0,"TotalContribution":daily_return-total_contrib})

        prev_nav=nav
        prev_values=security_values
        prev_prices=px

    ts=pd.DataFrame(rows)
    weights=pd.DataFrame(weight_rows)
    attribution=pd.DataFrame(attr_rows)
    if not ts.empty:
        ts["PortfolioGrowth"]=(1+pd.to_numeric(ts["PortfolioReturn"],errors="coerce").fillna(0)).cumprod()
        if adjusted_benchmark is not None and not adjusted_benchmark.empty:
            b=adjusted_benchmark.copy()
            b.index=pd.to_datetime(b.index).tz_localize(None)
            bret=b.pct_change().reindex(ts["Date"]).fillna(0).to_numpy()
            ts["BenchmarkReturn"]=bret
            ts["BenchmarkGrowth"]=(1+pd.Series(bret)).cumprod().to_numpy()
            ts["ActiveReturn"]=ts["PortfolioReturn"]-ts["BenchmarkReturn"]
    final_nav=float(ts.iloc[-1]["NAV"]) if not ts.empty else None
    twr=float(ts.iloc[-1]["PortfolioGrowth"]-1) if have_external and not ts.empty else None
    if have_external and final_nav is not None:
        external_cashflows.append((pd.Timestamp(ts.iloc[-1]["Date"]),final_nav))
    mwr=_xirr(external_cashflows) if have_external else None
    benchmark_twr=float(ts.iloc[-1]["BenchmarkGrowth"]-1) if "BenchmarkGrowth" in ts and not ts.empty else None
    days=max(1,(pd.Timestamp(ts.iloc[-1]["Date"])-pd.Timestamp(ts.iloc[0]["Date"])).days) if len(ts)>=2 else 1
    ann_twr=(1+twr)**(365.25/days)-1 if twr is not None and twr>-1 and days>=30 else None

    negative_cash=bool((ts["Cash"]<-1e-6).any()) if not ts.empty else False
    if negative_cash:
        warnings.append("Cash became negative at least once; check deposits, withdrawals, fees and trade execution history.")

    summary={
        "status":"PASS" if have_external and not negative_cash else "REVIEW",
        "has_external_cash_flows":bool(have_external),
        "start":ts.iloc[0]["Date"].date().isoformat() if not ts.empty else None,
        "end":ts.iloc[-1]["Date"].date().isoformat() if not ts.empty else None,
        "observations":int(len(ts)),
        "final_nav":final_nav,
        "twr_total":twr,
        "twr_annualized":ann_twr,
        "mwr_xirr":mwr,
        "benchmark_total_return":benchmark_twr,
        "active_total_return":twr-benchmark_twr if twr is not None and benchmark_twr is not None else None,
        "transaction_notional":transaction_notional,
        "turnover_vs_ending_nav":transaction_notional/final_nav if final_nav and final_nav>0 else None,
        "implementation_shortfall":implementation_shortfall if implementation_shortfall else None,
        "method_note":"TWR removes explicit DEPOSIT/WITHDRAW flows. MWR uses dated external flows plus ending NAV. Raw closes plus provider corporate actions are preferred for accounting.",
    }
    return {"summary":summary,"timeseries":ts,"weights":weights,"attribution":attribution,"transactions":events,"warnings":warnings}


def aggregate_realized_attribution(attribution: pd.DataFrame) -> pd.DataFrame:
    if attribution is None or attribution.empty:
        return pd.DataFrame()
    return (
        attribution.groupby("Ticker",as_index=False)[["PriceContribution","DividendContribution","TotalContribution"]]
        .sum()
        .sort_values("TotalContribution",ascending=False)
    )


def brinson_sector_attribution(
    point_in_time_weights: pd.DataFrame,
    raw_prices: pd.DataFrame,
    sector_map: dict[str,str],
    benchmark_sector_history_path: str | Path,
) -> pd.DataFrame:
    """Brinson allocation/selection/interaction using supplied point-in-time benchmark sector data.

    Input columns: Date,Sector,Weight,Return. No current-weight backfill is allowed.
    """
    path=Path(benchmark_sector_history_path)
    if not path.exists() or point_in_time_weights.empty:
        return pd.DataFrame()
    bench=pd.read_csv(path)
    required={"Date","Sector","Weight","Return"}
    if not required.issubset(set(bench.columns)):
        return pd.DataFrame()
    bench["Date"]=pd.to_datetime(bench["Date"],errors="coerce")
    bench["Weight"]=pd.to_numeric(bench["Weight"],errors="coerce")
    bench["Return"]=pd.to_numeric(bench["Return"],errors="coerce")
    bench=bench.dropna(subset=["Date","Sector","Weight","Return"])
    if bench.empty:
        return pd.DataFrame()

    weights=point_in_time_weights[point_in_time_weights["Ticker"]!="CASH"].copy()
    weights["Sector"]=weights["Ticker"].map(sector_map).fillna("Unknown")
    prices=raw_prices.copy(); prices.index=pd.to_datetime(prices.index).tz_localize(None)
    returns=prices.pct_change()
    rows=[]
    dates=sorted(set(weights["Date"]).intersection(set(bench["Date"])))
    prev_weight={}
    for dt in dates:
        dayw=weights[weights["Date"]==dt]
        bday=bench[bench["Date"]==dt]
        if dayw.empty or bday.empty:
            continue
        # Use start-of-period (previous available) portfolio sector weights when possible.
        current_sector=dayw.groupby("Sector")["Weight"].sum().to_dict()
        use_weight=prev_weight or current_sector
        if dt not in returns.index:
            prev_weight=current_sector; continue
        asset_day=returns.loc[dt]
        temp=dayw.copy()
        temp["AssetReturn"]=temp["Ticker"].map(asset_day.to_dict())
        sector_returns={}
        for sector,part in temp.groupby("Sector"):
            denom=part["Weight"].sum()
            sector_returns[sector]=float((part["Weight"]*part["AssetReturn"].fillna(0)).sum()/denom) if denom else 0.0
        bweight=dict(zip(bday["Sector"],bday["Weight"]))
        bret=dict(zip(bday["Sector"],bday["Return"]))
        total_bench=sum(bweight.get(s,0)*bret.get(s,0) for s in bweight)
        for sector in sorted(set(use_weight)|set(bweight)):
            wp=use_weight.get(sector,0.0); wb=bweight.get(sector,0.0)
            rp=sector_returns.get(sector,0.0); rb=bret.get(sector,0.0)
            alloc=(wp-wb)*(rb-total_bench)
            select=wb*(rp-rb)
            interact=(wp-wb)*(rp-rb)
            rows.append({"Date":dt,"Sector":sector,"Allocation":alloc,"Selection":select,"Interaction":interact,"TotalActiveContribution":alloc+select+interact})
        prev_weight=current_sector
    if not rows:
        return pd.DataFrame()
    return pd.DataFrame(rows)


def _price_on_or_after(series: pd.Series, date):
    s=series.dropna().copy(); s.index=pd.to_datetime(s.index).tz_localize(None)
    pos=s.index.searchsorted(pd.Timestamp(date),side="left")
    return (s.index[pos],float(s.iloc[pos])) if pos<len(s) else (None,None)


def _price_on_or_before(series: pd.Series, date):
    s=series.dropna().copy(); s.index=pd.to_datetime(s.index).tz_localize(None)
    pos=s.index.searchsorted(pd.Timestamp(date),side="right")-1
    return (s.index[pos],float(s.iloc[pos])) if pos>=0 else (None,None)


def decision_journal_analytics(path: str | Path, adjusted_prices: pd.DataFrame, benchmark: str) -> tuple[pd.DataFrame,pd.DataFrame]:
    path=Path(path)
    if not path.exists() or adjusted_prices is None or adjusted_prices.empty:
        return pd.DataFrame(),pd.DataFrame()
    df=pd.read_csv(path,comment="#")
    if df.empty or not {"Date","Ticker","Decision"}.issubset(df.columns):
        return pd.DataFrame(),pd.DataFrame()
    df["Date"]=pd.to_datetime(df["Date"],errors="coerce")
    df["Ticker"]=df["Ticker"].astype(str).str.upper().str.strip()
    df["Decision"]=df["Decision"].astype(str).str.upper()
    rows=[]
    for _,r in df.dropna(subset=["Date"]).iterrows():
        t=r["Ticker"]
        if t not in adjusted_prices.columns or benchmark not in adjusted_prices.columns:
            continue
        start_date,start_px=_price_on_or_after(adjusted_prices[t],r["Date"])
        _,start_b=_price_on_or_after(adjusted_prices[benchmark],r["Date"])
        if start_date is None or start_px in (None,0) or start_b in (None,0):
            continue
        for label,days in [("3M",91),("6M",182),("12M",365)]:
            end_target=start_date+pd.Timedelta(days=days)
            end_date,end_px=_price_on_or_before(adjusted_prices[t],end_target)
            _,end_b=_price_on_or_before(adjusted_prices[benchmark],end_target)
            matured=end_date is not None and end_date>=end_target-pd.Timedelta(days=7)
            if end_px is None or end_b is None:
                continue
            security=end_px/start_px-1; benchret=end_b/start_b-1; active=security-benchret
            decision=str(r["Decision"])
            sellish=("SELL" in decision or "TRIM" in decision)
            decision_alpha=-active if sellish else active
            rows.append({
                "Date":r["Date"],"Ticker":t,"Decision":r["Decision"],"Horizon":label,
                "StartDate":start_date,"EndDate":end_date,"Matured":bool(matured),
                "SecurityReturn":security,"BenchmarkReturn":benchret,"ActiveReturn":active,
                "DecisionAlpha":decision_alpha,"Correct":bool(decision_alpha>0) if matured else None,
                "ExpectedReturn":_num(r.get("ExpectedReturn")),
                "Conviction":_num(r.get("Conviction")),
                "PrimaryReason":r.get("PrimaryReason"),"KeyRisk":r.get("KeyRisk"),
                "WhatWouldChangeMyMind":r.get("WhatWouldChangeMyMind"),
            })
    detail=pd.DataFrame(rows)
    matured=detail[detail["Matured"]==True].copy() if not detail.empty else pd.DataFrame()
    if matured.empty:
        return detail,pd.DataFrame()
    summary=matured.groupby(["Horizon","Decision"],as_index=False).agg(
        Decisions=("Correct","count"),
        CorrectRate=("Correct","mean"),
        AverageDecisionAlpha=("DecisionAlpha","mean"),
        MedianDecisionAlpha=("DecisionAlpha","median"),
    )
    return detail,summary


def _find_row(ws, labels):
    if isinstance(labels,str): labels=[labels]
    needles={str(x).strip().lower() for x in labels}
    for r in range(1,ws.max_row+1):
        if str(ws.cell(r,1).value or "").strip().lower() in needles:
            return r
    return None


def find_latest_workbook(root: str | Path, ticker: str):
    root=Path(root)
    safe=re.sub(r"[^A-Za-z0-9]+","",ticker.upper())
    patterns=[f"{ticker.upper()}_Equity_Research_*.xlsx",f"{safe}_Equity_Research_*.xlsx"]
    candidates=[]
    for base in [root,root/"updated_models",root/"research_runs"]:
        if not base.exists():
            continue
        for pat in patterns:
            candidates.extend(base.glob(pat))
            candidates.extend(base.rglob(pat))
    candidates=[p for p in set(candidates) if "CLEAN" not in p.name.upper()]
    return max(candidates,key=lambda p:p.stat().st_mtime) if candidates else None


def _workbook_research_snapshot(path: Path):
    try:
        wb=load_workbook(path,data_only=True,read_only=True)
    except Exception:
        return {}
    out={"Workbook":str(path)}
    if "Company Data" in wb.sheetnames:
        ws=wb["Company Data"]; out["Price"]=_num(ws["B8"].value); out["Shares"]=_num(ws["B9"].value)
    if "Decision View" in wb.sheetnames:
        ws=wb["Decision View"]
        for label,key in [
            ("Base DCF Fair Value","BaseValue"),("Current Market Price","Price"),
            ("MODEL VIEW","ModelView"),("Monte Carlo P10","BearValue"),
            ("Overall Investment Score","OverallScore"),("Business Quality","QualityScore"),
            ("Valuation / Price","ValuationScore"),
        ]:
            r=_find_row(ws,label)
            if r:
                out[key]=ws.cell(r,2).value
    if "Advanced Analytics" in wb.sheetnames:
        ws=wb["Advanced Analytics"]
        for label,key in [("P10 Value / Share","BearValue"),("P90 Value / Share","BullValue"),("Median Value / Share","MedianValue"),("Probability > Current Price","MonteCarloProb")]:
            r=_find_row(ws,label)
            if r: out[key]=ws.cell(r,2).value
    if "Data Quality" in wb.sheetnames:
        ws=wb["Data Quality"]; p=review=fail=0
        for r in range(1,ws.max_row+1):
            s=str(ws.cell(r,2).value or "").upper().strip()
            p+=s=="PASS"; review+=s=="REVIEW"; fail+=s=="FAIL"
        total=p+review+fail
        out["DataQualityScore"]=(p+.5*review)/total if total else None
        out["DataQualityFail"]=fail
    if "Forecast Accountability" in wb.sheetnames:
        ws=wb["Forecast Accountability"]
        by_metric={"FCF":[],"Revenue":[]}
        for r in range(7,min(ws.max_row,40)+1):
            metric=str(ws.cell(r,2).value or "").strip()
            if metric not in by_metric:
                continue
            year=_num(ws.cell(r,1).value); value=_num(ws.cell(r,3).value)
            if year and value is not None and value>0:
                by_metric[metric].append((int(year),value))
        forecasts=by_metric["FCF"] if len(by_metric["FCF"])>=2 else by_metric["Revenue"]
        if len(forecasts)>=2 and forecasts[0][1]>0:
            y0,v0=forecasts[0]; y1,v1=forecasts[-1]
            out["BaseFCFCAGRProxy"]=(v1/v0)**(1/max(1,y1-y0))-1
            out["BaseGrowthProxySource"]="FCF" if len(by_metric["FCF"])>=2 else "Revenue fallback"
    if "Capital Allocation" in wb.sheetnames:
        ws=wb["Capital Allocation"]
        r=_find_row(ws,"Latest net share reduction")
        if r: out["NetBuybackYield"]=_num(ws.cell(r,2).value)
    return out


def _forecast_track_record(root: str | Path, ticker: str) -> dict:
    path=Path(root)/"research_data"/str(ticker).upper()/"forecast_accuracy_summary.csv"
    if not path.exists():
        return {"score":0.50,"mae":None,"observations":0,"status":"INSUFFICIENT_HISTORY"}
    try:
        df=pd.read_csv(path)
    except Exception:
        return {"score":0.50,"mae":None,"observations":0,"status":"REVIEW"}
    if df.empty or "MeanAbsoluteError" not in df.columns:
        return {"score":0.50,"mae":None,"observations":0,"status":"INSUFFICIENT_HISTORY"}
    errors=pd.to_numeric(df["MeanAbsoluteError"],errors="coerce").dropna()
    obs_series=pd.to_numeric(df.get("Observations",pd.Series(1,index=df.index)),errors="coerce").fillna(1)
    observations=int(obs_series.sum())
    if errors.empty:
        return {"score":0.50,"mae":None,"observations":observations,"status":"INSUFFICIENT_HISTORY"}
    weights=np.maximum(1,obs_series.reindex(errors.index).to_numpy(float))
    mae=float(np.average(errors.to_numpy(float),weights=weights))
    score=float(max(.20,min(1.0,1-mae/.25)))
    return {"score":score,"mae":mae,"observations":observations,"status":"PASS" if observations>=3 else "LIMITED_HISTORY"}


def research_expected_return_bridge(
    tickers: list[str],
    root: str | Path,
    info: dict[str,dict],
    benchmark_expected_return: float=0.08,
    convergence_years: float=3.0,
) -> pd.DataFrame:
    rows=[]
    for ticker in tickers:
        path=find_latest_workbook(root,ticker)
        snap=_workbook_research_snapshot(path) if path else {}
        price=_num(snap.get("Price"))
        base=_num(snap.get("BaseValue"))
        bear=_num(snap.get("BearValue"))
        bull=_num(snap.get("BullValue"))
        dividend=_num((info.get(ticker) or {}).get("dividendYield"),0.0) or 0.0
        raw=None
        if price and base and price>0 and base>0:
            raw=(base/price)**(1/max(.25,float(convergence_years)))-1+dividend
        quality=_num(snap.get("DataQualityScore"),0.45)
        if _num(snap.get("DataQualityFail"),0)>0:
            quality=min(quality,0.25)
        score=_num(snap.get("OverallScore"))
        score_conf=0.5 if score is None else min(1,max(0,score/100))
        track=_forecast_track_record(root,ticker)
        confidence=min(1,max(.15,.55*quality+.20*score_conf+.25*track["score"]))
        adjusted=benchmark_expected_return+confidence*(raw-benchmark_expected_return) if raw is not None else None
        rows.append({
            "Ticker":ticker,"Workbook":str(path) if path else None,
            "CurrentPrice":price,"BearValue":bear,"BaseValue":base,"BullValue":bull,
            "RawExpectedReturn":raw,"Confidence":confidence if raw is not None else 0.0,
            "ConfidenceAdjustedExpectedReturn":adjusted,
            "ExpectedAlpha":adjusted-benchmark_expected_return if adjusted is not None else None,
            "DividendYield":dividend,
            "BaseGrowthProxy":_num(snap.get("BaseFCFCAGRProxy")),
            "NetBuybackYield":_num(snap.get("NetBuybackYield")),
            "BaseGrowthProxySource":snap.get("BaseGrowthProxySource"),
            "DataQualityScore":quality if path else None,
            "ForecastAccuracyScore":track["score"] if path else None,
            "ForecastMAE":track["mae"] if path else None,
            "ForecastAccuracyObservations":track["observations"] if path else 0,
            "ForecastTrackRecordStatus":track["status"] if path else "NO_WORKBOOK",
            "ModelView":snap.get("ModelView"),
            "SourceStatus":"PASS" if path and raw is not None else "REVIEW",
        })
    return pd.DataFrame(rows)


def position_sizing_ranges(
    holdings: pd.DataFrame,
    research_bridge: pd.DataFrame,
    max_position: float,
    half_width: float=0.025,
    liquidity: pd.DataFrame | None=None,
    max_days_to_liquidate: float=5.0,
) -> pd.DataFrame:
    if holdings.empty:
        return pd.DataFrame()
    h=holdings.copy().set_index("Ticker")
    r=research_bridge.set_index("Ticker") if research_bridge is not None and not research_bridge.empty else pd.DataFrame()
    liq=liquidity.set_index("Ticker") if liquidity is not None and not liquidity.empty and "Ticker" in liquidity else pd.DataFrame()
    rows=[]
    for ticker,row in h.iterrows():
        current=_num(row.get("Weight"),0.0) or 0.0
        risk=_num(row.get("RiskContributionPct"),current) or current
        if not r.empty and ticker in r.index:
            rr=r.loc[ticker]
            alpha=_num(rr.get("ExpectedAlpha"))
            conf=_num(rr.get("Confidence"),0.0) or 0.0
            price=_num(rr.get("CurrentPrice")); bear=_num(rr.get("BearValue"))
            downside=abs(min(0,bear/price-1)) if price and bear else .30
        else:
            alpha=None; conf=0.0; downside=.30
        opportunity=0.0 if alpha is None else max(0,min(1,(alpha+.02)/.14))
        downside_penalty=max(.35,1-min(.65,downside))
        liquidity_days=None
        liquidity_factor=1.0
        if not liq.empty and ticker in liq.index and "EstimatedDaysToLiquidate" in liq.columns:
            liquidity_days=_num(liq.loc[ticker,"EstimatedDaysToLiquidate"])
            if liquidity_days is not None and liquidity_days>max_days_to_liquidate:
                liquidity_factor=max(.40,min(1.0,max_days_to_liquidate/liquidity_days))
        target=max_position*opportunity*conf*downside_penalty*liquidity_factor
        # RiskContributionPct is covariance-based, so this guardrail incorporates the
        # holding's volatility and correlation with the rest of the portfolio.
        if current>0 and risk/current>1.35:
            target*=0.85
        lower=max(0,target-half_width)
        upper=min(max_position,target+half_width)
        status="WITHIN RANGE" if lower-1e-9<=current<=upper+1e-9 else "ABOVE RANGE" if current>upper else "BELOW RANGE"
        rows.append({
            "Ticker":ticker,"CurrentWeight":current,"RiskContribution":risk,
            "ExpectedAlpha":alpha,"ResearchConfidence":conf,"DownsideReference":downside,
            "LiquidityDays":liquidity_days,"LiquidityFactor":liquidity_factor,
            "SuggestedMidpoint":target,"SuggestedMin":lower,"SuggestedMax":upper,"RangeStatus":status,
            "Method":"confidence-adjusted alpha × downside × liquidity × covariance-based risk guardrail",
        })
    return pd.DataFrame(rows)


def thesis_budget(holdings: pd.DataFrame, research_bridge: pd.DataFrame) -> pd.DataFrame:
    if holdings.empty:
        return pd.DataFrame()
    h=holdings[["Ticker","Weight"]].copy()
    if "RiskContributionPct" in holdings:
        h["RiskWeight"]=holdings["RiskContributionPct"].values
    else:
        h["RiskWeight"]=h["Weight"]
    r=research_bridge[["Ticker","ExpectedAlpha","ConfidenceAdjustedExpectedReturn","Confidence"]].copy() if research_bridge is not None and not research_bridge.empty else pd.DataFrame(columns=["Ticker","ExpectedAlpha","ConfidenceAdjustedExpectedReturn","Confidence"])
    out=h.merge(r,on="Ticker",how="left")
    positive=pd.to_numeric(out["ExpectedAlpha"],errors="coerce").clip(lower=0).fillna(0)*pd.to_numeric(out["Confidence"],errors="coerce").fillna(0)
    total=positive.sum()
    out["ExpectedAlphaBudget"]=positive/total if total>0 else np.nan
    out["RiskToCapital"]=out["RiskWeight"]/out["Weight"].replace(0,np.nan)
    out["AlphaToCapital"]=out["ExpectedAlphaBudget"]/out["Weight"].replace(0,np.nan)
    return out


def research_fundamental_scenarios(holdings: pd.DataFrame, bridge: pd.DataFrame) -> pd.DataFrame:
    if holdings.empty or bridge is None or bridge.empty:
        return pd.DataFrame()
    df=holdings[["Ticker","Weight"]].merge(bridge,on="Ticker",how="left")
    rows=[]
    for scenario,col in [("Research Bear","BearValue"),("Research Base","BaseValue"),("Research Bull / P90","BullValue")]:
        contributions=[]; covered=0.0
        for _,r in df.iterrows():
            price=_num(r.get("CurrentPrice")); value=_num(r.get(col)); w=_num(r.get("Weight"),0.0) or 0.0
            shock=value/price-1 if price and value else None
            if shock is not None:
                covered+=w; contributions.append((r["Ticker"],w,shock,w*shock))
        if not contributions:
            continue
        portfolio=sum(x[3] for x in contributions)
        rows.append({"Scenario":scenario,"PortfolioShock":portfolio,"CoveredWeight":covered,"UncoveredWeight":max(0,1-covered),"Method":"company-model valuation / current price"})
        for t,w,shock,contrib in contributions:
            rows.append({"Scenario":scenario,"Ticker":t,"Weight":w,"SecurityShock":shock,"Contribution":contrib,"CoveredWeight":covered,"Method":"company-model valuation / current price"})
    return pd.DataFrame(rows)


def custom_thesis_scenarios(path: str | Path, holdings: pd.DataFrame) -> pd.DataFrame:
    """Aggregate explicit company-level thesis shocks; shocks are never inferred from beta."""
    path=Path(path)
    if not path.exists() or holdings is None or holdings.empty:
        return pd.DataFrame()
    try:
        raw=pd.read_csv(path,comment="#")
    except Exception:
        return pd.DataFrame()
    required={"Scenario","Ticker","Shock"}
    if raw.empty or not required.issubset(raw.columns):
        return pd.DataFrame()
    raw["Ticker"]=raw["Ticker"].astype(str).str.upper().str.strip()
    raw["Shock"]=pd.to_numeric(raw["Shock"],errors="coerce")
    weights=holdings.set_index("Ticker")["Weight"].to_dict()
    rows=[]
    for scenario,part in raw.dropna(subset=["Shock"]).groupby("Scenario"):
        covered=0.0; total=0.0; details=[]
        for _,row in part.iterrows():
            t=row["Ticker"]
            if t not in weights:
                continue
            w=float(weights[t]); shock=float(row["Shock"]); contrib=w*shock
            covered+=w; total+=contrib
            details.append({
                "Scenario":scenario,"Ticker":t,"Weight":w,"SecurityShock":shock,
                "Contribution":contrib,"Notes":row.get("Notes"),
                "Method":"explicit local thesis-scenario input",
            })
        if not details:
            continue
        rows.append({
            "Scenario":scenario,"PortfolioShock":total,"CoveredWeight":covered,
            "UncoveredWeight":max(0.0,1-covered),"Method":"explicit local thesis-scenario input",
        })
        rows.extend(details)
    return pd.DataFrame(rows)


def expected_return_decomposition(holdings: pd.DataFrame, bridge: pd.DataFrame) -> tuple[pd.DataFrame,pd.DataFrame]:
    if holdings.empty or bridge is None or bridge.empty:
        return pd.DataFrame(),pd.DataFrame()
    df=holdings[["Ticker","Weight"]].merge(bridge,on="Ticker",how="left")
    rows=[]
    for _,r in df.iterrows():
        total=_num(r.get("ConfidenceAdjustedExpectedReturn"))
        growth=_num(r.get("BaseGrowthProxy"),0.0) or 0.0
        div=_num(r.get("DividendYield"),0.0) or 0.0
        buyback=_num(r.get("NetBuybackYield"),0.0) or 0.0
        residual=total-growth-div-buyback if total is not None else None
        rows.append({"Ticker":r["Ticker"],"Weight":r["Weight"],"ExpectedReturn":total,"FundamentalGrowth":growth,"DividendYield":div,"NetBuybackYield":buyback,"ValuationNormalizationResidual":residual})
    detail=pd.DataFrame(rows)
    summary=[]
    for col,label in [("FundamentalGrowth","Expected fundamental growth"),("DividendYield","Dividend yield"),("NetBuybackYield","Net buyback yield"),("ValuationNormalizationResidual","Valuation / other residual"),("ExpectedReturn","Total confidence-adjusted expected return")]:
        valid=detail.dropna(subset=[col])
        value=float((valid["Weight"]*valid[col]).sum()) if not valid.empty else None
        summary.append({"Component":label,"PortfolioContribution":value})
    return detail,pd.DataFrame(summary)


def transaction_cost_rebalance(
    optimizer_weights: pd.DataFrame,
    holdings: pd.DataFrame,
    liquidity: pd.DataFrame,
    research_bridge: pd.DataFrame,
    portfolio_name: str="Regime-Aware ML Maximum Sharpe",
    rebalance_threshold: float=0.03,
    spread_bps: float=5.0,
    impact_bps_at_10pct_adv: float=20.0,
    min_benefit_cost_ratio: float=1.5,
) -> pd.DataFrame:
    if optimizer_weights is None or optimizer_weights.empty or holdings.empty:
        return pd.DataFrame()
    target=optimizer_weights[optimizer_weights["Portfolio"]==portfolio_name].copy()
    if target.empty:
        # Prefer an expected-return aware portfolio rather than silently using minimum variance.
        names=optimizer_weights["Portfolio"].astype(str)
        cand=optimizer_weights[names.str.contains("ML Maximum Sharpe",case=False,regex=False)]
        target=cand.copy() if not cand.empty else pd.DataFrame()
    if target.empty:
        return pd.DataFrame()
    market_value=float(pd.to_numeric(holdings.get("MarketValue"),errors="coerce").sum())
    if not market_value or market_value<=0:
        return pd.DataFrame()
    liq=liquidity.set_index("Ticker") if liquidity is not None and not liquidity.empty and "Ticker" in liquidity else pd.DataFrame()
    rb=research_bridge.set_index("Ticker") if research_bridge is not None and not research_bridge.empty else pd.DataFrame()
    rows=[]
    for _,r in target.iterrows():
        t=r["Ticker"]; change=_num(r.get("WeightChange"),0.0) or 0.0
        trade_value=abs(change)*market_value
        adv=None
        if not liq.empty and t in liq.index:
            for key in ["AverageDailyDollarVolume","ADV_Dollars","DollarADV","AverageDailyValue"]:
                if key in liq.columns:
                    adv=_num(liq.loc[t,key]); 
                    if adv: break
        participation=trade_value/adv if adv else None
        impact=impact_bps_at_10pct_adv*math.sqrt(max(0,participation)/.10) if participation is not None else impact_bps_at_10pct_adv
        cost_bps=spread_bps+impact
        alpha=_num(rb.loc[t,"ExpectedAlpha"]) if not rb.empty and t in rb.index else None
        annual_benefit_bps=max(0,alpha or 0)*abs(change)*10000
        ratio=annual_benefit_bps/cost_bps if cost_bps>0 else None
        action="REBALANCE" if abs(change)>=rebalance_threshold and ratio is not None and ratio>=min_benefit_cost_ratio else "HOLD / NO TRADE"
        rows.append({
            "Ticker":t,"CurrentWeight":r.get("CurrentWeight"),"TargetWeight":r.get("TargetWeight"),"WeightChange":change,
            "EstimatedTradeValue":trade_value,"ADV":adv,"TradeAsPctADV":participation,
            "EstimatedCostBps":cost_bps,"AnnualExpectedBenefitBps":annual_benefit_bps,
            "BenefitCostRatio":ratio,"Decision":action,
            "CostMethod":"configured spread + square-root ADV impact diagnostic",
        })
    return pd.DataFrame(rows)


def high_level_decision_learning(decision_detail: pd.DataFrame) -> pd.DataFrame:
    if decision_detail is None or decision_detail.empty:
        return pd.DataFrame()
    x=decision_detail[decision_detail["Matured"]==True].copy()
    if x.empty:
        return pd.DataFrame()
    x["ConvictionBucket"]=pd.cut(pd.to_numeric(x["Conviction"],errors="coerce"),bins=[-np.inf,2.5,3.5,np.inf],labels=["Low","Medium","High"])
    return x.groupby(["Horizon","ConvictionBucket"],observed=True,as_index=False).agg(
        Decisions=("DecisionAlpha","count"),
        CorrectRate=("Correct","mean"),
        AverageDecisionAlpha=("DecisionAlpha","mean"),
        MedianDecisionAlpha=("DecisionAlpha","median"),
    )
